"""Two tweaks to the results slide:
  1) replace whole-word 'his' -> 'prior' (neutral wording, no finger-pointing)
  2) append a note: the prior gaze QDA is ITSELF n.s. vs chance after Bonferroni
     (medium-large effects but n=8-20 underpowered) -> nobody clears the bar.
Idempotent. Run from repo root.
"""
import re
from pptx import Presentation
from pptx.util import Pt

SRC = "discover/2026_March_ViLearn_Planning.pptx"
prs = Presentation(SRC)
HIS = re.compile(r"\bhis\b", re.I)

# 1) neutralize 'his' across all table cells + text frames
for s in prs.slides:
    for sh in s.shapes:
        if sh.has_table:
            for r in range(len(sh.table.rows)):
                for c in range(len(sh.table.columns)):
                    cell = sh.table.cell(r, c)
                    if HIS.search(cell.text):
                        new = HIS.sub("prior", cell.text)
                        cell.text = new
                        run = cell.text_frame.paragraphs[0].runs[0]
                        run.font.size = Pt(12)
        elif sh.has_text_frame and HIS.search(sh.text_frame.text):
            for p in sh.text_frame.paragraphs:
                for run in p.runs:
                    if HIS.search(run.text):
                        run.text = HIS.sub("prior", run.text)

# 2) results-slide caption: set full text incl. prior-vs-chance note
NEW_CAPTION = (
    "Per-fold LOSO accuracy; paired t-test (Bonferroni) + Cohen's d. Harmonized 2-annotator "
    "labels + identical nested CV. Our detector and the prior gaze QDA are statistically "
    "indistinguishable on every split (all n.s.) — we edge triads, the prior edges dyads. "
    "The prior QDA is itself n.s. vs chance after Bonferroni too (d up to 1.18, but n=8–20 "
    "underpowered): nothing here — ours or prior — survives correction. The earlier "
    "'beats prior' result was a label-distribution artifact."
)
for s in prs.slides:
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().startswith("Per-fold"):
            tf = sh.text_frame
            for p in tf.paragraphs[1:]:
                p._p.getparent().remove(p._p)
            p0 = tf.paragraphs[0]
            for r in list(p0.runs)[1:]:
                r._r.getparent().remove(r._r)
            if p0.runs:
                p0.runs[0].text = NEW_CAPTION
            else:
                p0.text = NEW_CAPTION

prs.save(SRC)
print("neutralized 'his' -> 'prior'; added prior-vs-chance note to results caption")
