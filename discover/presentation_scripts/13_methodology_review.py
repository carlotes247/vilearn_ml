"""Append a 'Methodology Review (for next meeting)' section: 3 slides documenting
the decisions that move the significance verdict, so the team can pick a method.
Idempotent. Run from repo root, after the other builders."""
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt

SRC = "discover/2026_March_ViLearn_Planning.pptx"
prs = Presentation(SRC)
mt = pd.read_csv("scratch/te_adoption/a1_out/methods_table.csv")


def has_slide(title):
    return any(any(sh.has_text_frame and sh.text_frame.text.strip() == title for sh in s.shapes)
               for s in prs.slides)


def add_title_only(title):
    s = prs.slides.add_slide(prs.slide_layouts[7])
    s.shapes.title.text = title
    return s


def add_table(slide, header, rows, top, widths, fontsz=11):
    tbl = slide.shapes.add_table(len(rows) + 1, len(header), Inches(0.4), Inches(top),
                                 Inches(sum(widths)), Inches(0.36 * (len(rows) + 1))).table
    for j, wdt in enumerate(widths):
        tbl.columns[j].width = Inches(wdt)
    for j, h in enumerate(header):
        c = tbl.cell(0, j); c.text = h
        c.text_frame.paragraphs[0].runs[0].font.size = Pt(fontsz)
        c.text_frame.paragraphs[0].runs[0].font.bold = True
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row):
            c = tbl.cell(i, j); c.text = str(v)
            c.text_frame.paragraphs[0].runs[0].font.size = Pt(fontsz)
    return tbl


def note(slide, text, top, h=1.4):
    tb = slide.shapes.add_textbox(Inches(0.4), Inches(top), Inches(12.4), Inches(h))
    tb.text_frame.word_wrap = True
    for i, line in enumerate(text):
        p = tb.text_frame.paragraphs[0] if i == 0 else tb.text_frame.add_paragraph()
        p.text = line
        for r in p.runs:
            r.font.size = Pt(12)


def sigmark(p):
    return f"{p:.3f} {'✓' if p < 0.05 else '✗'}"


# ---- Slide A: significance method comparison (the flip-flop) ----
if not has_slide("Significance Depends on the Test (sig → ns → sig)"):
    s = add_title_only("Significance Depends on the Test (sig → ns → sig)")
    rows = []
    for _, r in mt.iterrows():
        rows.append([f"{r.detector} {r.split}", f"{r.best_model} {r.acc:.2f}", f"{r.majority:.2f}",
                     sigmark(r.p_uniform_paired), sigmark(r.p_majority_1samp), sigmark(r.p_permutation)])
    add_table(s, ["Detector / split", "best acc", "majority",
                  "uniform-dummy", "majority 1-samp", "permutation"], rows,
              top=1.4, widths=[2.6, 2.4, 1.4, 2.3, 2.3, 2.0])
    note(s, [
        "All group-level, nested model selection, RAW p (Bonferroni shifts the threshold but not the pattern).",
        "Uniform-dummy = stochastic/gameable floor (prior + our first pass). Majority = deterministic always-predict-majority. Permutation = shuffle-label null (seed-proof).",
        "Verdict flips by method: audio-Triads & gaze-Triads are significant ONLY under permutation; gaze-Dyads LOSES significance under permutation (6-feature null fits shuffled labels). Decision needed: which test do we standardize on?",
    ], top=4.2)

# ---- Slide B: representation (group-level vs long-format) ----
if not has_slide("Representation: Group-level vs Long-format"):
    s = add_title_only("Representation: Group-level vs Long-format")
    add_table(s, ["Audio (best model)", "group-level (197)", "long-format (506)", "Δ"],
              [["All", "0.68", "0.68", "+0.01"], ["Dyads", "0.72", "0.71", "-0.01"],
               ["Triads", "0.70", "0.70", "0.00"]],
              top=1.6, widths=[3.0, 3.0, 3.0, 1.5])
    note(s, [
        "Long-format repeats each group-window once per participant (dyad ×2, triad ×3) → pseudo-replication on a group-level label.",
        "But best-model accuracy is ~identical (Δ ≤ 0.01) → dedup does NOT inflate numbers. Group-level is the honest unit and avoids the reviewer objection.",
        "Note: the prior gaze detector was ALWAYS group-level (185 windows). Moving ours to group-level makes the comparison apples-to-apples.",
    ], top=3.6)

# ---- Slide C: open decisions ----
if not has_slide("Open Methodology Decisions"):
    s = add_title_only("Open Methodology Decisions")
    note(s, [
        "1. Significance test: uniform-dummy (gameable) vs majority one-sample vs permutation (recommend permutation — seed-proof, handles class balance).",
        "2. Unit of analysis: group-level (honest, recommend) vs long-format (pseudo-replication). Accuracy ~unchanged either way.",
        "3. Chance baseline: uniform 0.5 vs deterministic majority (All 0.51 / Dyads 0.59 / Triads 0.55). Recommend majority.",
        "4. Model selection: nested best-per-cell — the dyad verdict flips on which model is used (single fixed model is misleading).",
        "5. Headline modality: audio-only (best for group TE, gaze-free) vs multimodal. Recommend lead with audio-only.",
        "6. Dyads n=8 is underpowered — report with caveat regardless of method.",
        "7. Scaled vs unscaled features (slide 'Prior Baseline — Corrected'): scaled = our pipeline; affects prior 0.67→0.73.",
        "8. TODO: recompute the head-to-head (ours vs prior) at group-level + permutation for a final fair comparison.",
    ], top=1.4, h=5.0)

prs.save(SRC)
print(f"methodology-review slides added; {len(prs.slides)} slides total")
