"""Linguistic-v2 in-place deck patch (2026-07-07, saveli's feature-naming decision).

Final linguistic set = 7 raw-ASR-segment stats: segments/window, words/window,
avg word length, question ratio, speech ratio, mean segment duration, and NEW
unfinished ratio (ends '...' OR no .!? terminal and not same-speaker-continued <2s;
trail-off / abandoned speech — cognitive-load / interruption marker).

IN-PLACE ONLY: cell texts + textbox paragraphs are rewritten, one row-triple inserted
into slide 21's table (XML copy). NO shapes are moved/rebuilt — manual layout survives.

  Slide 11  + one sentence naming the paper linguistic set
  Slide 13  linguistic + fusion rows: new numbers
  Slide 15  best-single/best-fusion rows recomputed + renamed
  Slide 17  feature rows renamed + new d1 numbers + refreshed bullets
  Slide 18  linguistic rows: new numbers
  Slide 21  + linguistic row per split (inserted after 'audio')
  Slide 23  linguistic rows: new numbers + label

Run from repo root:  python3 discover/presentation_scripts/24_linguistic_v2_inplace.py
"""
from copy import deepcopy

import pandas as pd
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Pt

SRC = "scratch/2026_March_ViLearn_Planning.pptx"
prs = Presentation(SRC)
c2 = pd.read_csv("scratch/te_adoption/a1_out/c2_paper_table.csv")
d1 = pd.read_csv("scratch/te_adoption/a1_out/d1_linguistic_dropone.csv")

SPL = {"Overall": "All", "Dyads": "D", "Triads": "T"}
LING = ["segment_count", "word_count", "avg_word_length", "question_ratio",
        "speech_ratio", "mean_segment_duration_s", "unfinished_ratio"]
DISPLAY = {"segment_count": "segments / window", "word_count": "words / window",
           "avg_word_length": "avg word length", "question_ratio": "question ratio",
           "speech_ratio": "speech ratio", "mean_segment_duration_s": "mean segment duration (s)",
           "unfinished_ratio": "unfinished ratio (NEW)"}


def pfmt(p):
    if pd.isna(p) or p == "":
        return "—"
    p = float(p)
    return "<.001" if p < 0.001 else f"{p:.3f}"


def r1(sp, det):
    m = c2[(c2.split == sp) & (c2.detector == det) & (c2["rank"] == 1)]
    return m.iloc[0] if len(m) else None


def set_text(cell, text, size=None, bold=None):
    old = cell.text_frame.paragraphs[0].runs
    sz = old[0].font.size if old else Pt(9)
    bd = old[0].font.bold if old else False
    cell.text = str(text)
    r = cell.text_frame.paragraphs[0].runs[0]
    r.font.size = size or sz
    r.font.bold = bd if bold is None else bold


def stats_cells(row, r):
    """Write model/acc/d/p-perm/p-t into columns 2..6 of a table row."""
    vals = [r.model, f"{r.acc_mean:.3f} ± {r.acc_std:.3f}", f"{r.cohen_d_maj:.2f}",
            pfmt(r.perm_p), pfmt(r.p_ttest_maj)]
    for j, v in enumerate(vals, start=2):
        set_text(row.cells[j], v)


def the_table(slide):
    return next(sh.table for sh in slide.shapes if sh.has_table)


def find_slide(prefix):
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip().startswith(prefix):
                return s
    return None


def update_rows(slide, mapping, rename=None):
    """mapping: {row-label -> detector}; updates stats cells of matching rows."""
    tbl = the_table(slide)
    n = 0
    for row in tbl.rows:
        sp_lab = row.cells[0].text.strip()
        lab = row.cells[1].text.strip()
        if sp_lab in SPL and lab in mapping:
            r = r1(SPL[sp_lab], mapping[lab])
            if r is not None:
                stats_cells(row, r)
                if rename and lab in rename:
                    set_text(row.cells[1], rename[lab])
                n += 1
    return n


# ---------- slide 11: name the paper set ----------
s11 = find_slide("Feature Modalities")
MARK = "Paper linguistic set"
for sh in s11.shapes:
    if sh.has_text_frame and "Ablation:" in sh.text_frame.text and MARK not in sh.text_frame.text:
        p = sh.text_frame.add_paragraph()
        p.text = (f"{MARK} (group-window level, raw ASR segments): segments/window, words/window, "
                  "avg word length, question ratio, speech ratio, mean segment duration, "
                  "unfinished ratio (trail-off '…' or abandoned segment — load/interruption marker).")
        for run in p.runs:
            run.font.size = Pt(11)
            run.font.bold = True
        print("slide 11: paper-set sentence appended")

# ---------- slides 13 / 18 / 23: refresh linguistic rows ----------
n13 = update_rows(find_slide("Modality Ablation"), {
    "linguistic (transcript stats)": "linguistic",
    "audio+linguistic": "audio+linguistic",
    "multimodal (audio+ling+gaze)": "audio+linguistic+GazexSpeaking"})
print(f"slide 13: {n13} rows refreshed")

n18 = update_rows(find_slide("Do Audio / Text Embeddings Help"),
                  {"handcrafted: linguistic stats": "linguistic"})
print(f"slide 18: {n18} rows refreshed")

n23 = update_rows(find_slide("Extended Feature Sets"),
                  {"linguistic (7 transcript stats)": "linguistic"},
                  rename={"linguistic (7 transcript stats)": "linguistic (7 segment stats)"})
print(f"slide 23: {n23} rows refreshed")

# ---------- slide 15: recompute best single / best fusion ----------
s15 = find_slide("Main Results — Group-Level")
tbl = the_table(s15)
singles = ["AIxVR", "GazexSpeaking", "audio", "linguistic"]
fusions = ["audio+AIxVR", "audio+GazexSpeaking", "audio+linguistic",
           "linguistic+GazexSpeaking", "audio+linguistic+GazexSpeaking"]
for row in tbl.rows:
    sp_lab = row.cells[0].text.strip()
    lab = row.cells[1].text.strip()
    if sp_lab not in SPL:
        continue
    sp = SPL[sp_lab]
    if lab.startswith("best single:"):
        b = max((r1(sp, d) for d in singles if r1(sp, d) is not None), key=lambda r: r.acc_mean)
        set_text(row.cells[1], f"best single: {b.detector}")
        stats_cells(row, b)
    elif lab.startswith("best fusion:"):
        b = max((r1(sp, d) for d in fusions if r1(sp, d) is not None), key=lambda r: r.acc_mean)
        set_text(row.cells[1], f"best fusion: {b.detector}")
        stats_cells(row, b)
print("slide 15: best-single/best-fusion refreshed")

# ---------- slide 17: rename features + new d1 numbers + bullets ----------
s17 = find_slide("What Drives the Linguistic Detector")
tbl = the_table(s17)
from sklearn.linear_model import LogisticRegression
from sklearn.preprocessing import StandardScaler
feats_df = pd.read_csv("scratch/te_adoption/a1_out/c2_features.csv")
coefs = {}
for sp in ("All", "D", "T"):
    sub = feats_df if sp == "All" else feats_df[feats_df.grp == sp]
    X = StandardScaler().fit_transform(sub[LING].values)
    lr = LogisticRegression(solver="liblinear", random_state=42).fit(X, sub["y"].values)
    coefs[sp] = dict(zip(LING, lr.coef_[0]))
for i, f in enumerate(LING, 1):
    solo = d1[(d1.split == "All") & (d1.variant == f"solo_{f}")].iloc[0]
    drop = d1[(d1.split == "All") & (d1.variant == f"drop_{f}")].iloc[0]
    row = list(tbl.rows)[i]
    vals = [DISPLAY[f], f"{solo.acc_mean:.3f}", f"{drop.delta_vs_full:+.3f}",
            *(f"{coefs[sp][f]:+.2f}" for sp in ("All", "D", "T"))]
    for j, v in enumerate(vals):
        set_text(row.cells[j], v)
full_all = d1[(d1.split == "All") & (d1.variant == "full")].iloc[0]
best_solo = d1[(d1.split == "All") & d1.variant.str.startswith("solo_")].sort_values("acc_mean").iloc[-1]
un_solo = d1[(d1.split == "All") & (d1.variant == "solo_unfinished_ratio")].iloc[0]
un_drop = d1[(d1.split == "All") & (d1.variant == "drop_unfinished_ratio")].iloc[0]
bullets = [
    f"Final 7-feature set (raw ASR segments; words_per_second removed as an exact word-count duplicate). "
    f"Full set Overall: {full_all.acc_mean:.3f} ± {full_all.acc_std:.3f}.",
    f"Signal stays DISTRIBUTED: best solo = {DISPLAY.get(best_solo.variant[5:], best_solo.variant[5:])} "
    f"{best_solo.acc_mean:.3f}; no drop-one collapses the score.",
    "question ratio keeps its NEGATIVE sign (more questions → low TE: confusion / clarification-seeking).",
    f"NEW unfinished ratio (trail-off '…' / abandoned segments — cognitive-load or interruption marker): "
    f"solo {un_solo.acc_mean:.3f}, drop-one Δ {un_drop.delta_vs_full:+.3f}, coef All {coefs['All']['unfinished_ratio']:+.2f}.",
    "Segments are raw ASR chunks (Whisper-style), NOT sentences: median 1.7 s, 66% end in terminal "
    "punctuation. Whisper strips fillers, so classic disfluency counts are unavailable.",
]
for sh in s17.shapes:
    if sh.has_text_frame and "DISTRIBUTED" in sh.text_frame.text:
        tf = sh.text_frame
        tf.clear()
        for i, t in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = t
            for run in p.runs:
                run.font.size = Pt(10)
        print("slide 17: bullets rewritten")
        break
print("slide 17: table renamed + refreshed")

# ---------- slide 21: insert linguistic row per split ----------
s21 = find_slide("Detector vs 3 Priors")
tbl = the_table(s21)
if not any(r.cells[1].text.strip() == "linguistic" for r in tbl.rows):
    tr_list = tbl._tbl.findall(qn("a:tr"))
    # find 'audio' row index per split (rows offset by header at 0)
    inserts = []
    for i, row in enumerate(tbl.rows):
        if row.cells[1].text.strip() == "audio":
            inserts.append((i, row.cells[0].text.strip()))
    for off, (i, sp_lab) in enumerate(inserts):
        src = tr_list[i]
        new = deepcopy(src)
        src.addnext(new)
    tbl = the_table(s21)  # re-read after XML change
    # rewrite the duplicated rows (they currently say 'audio' twice per split)
    seen = set()
    for row in tbl.rows:
        key = (row.cells[0].text.strip(), row.cells[1].text.strip())
        if key[1] == "audio":
            if key in seen:  # the duplicate -> becomes linguistic
                sp = SPL[key[0]]
                set_text(row.cells[1], "linguistic")
                stats_cells(row, r1(sp, "linguistic"))
            else:
                seen.add(key)
    print("slide 21: linguistic rows inserted (table is 3 rows taller — check layout)")
else:
    print("slide 21: linguistic rows already present")

prs.save(SRC)
print(f"saved {SRC}: {len(prs.slides)} slides")
