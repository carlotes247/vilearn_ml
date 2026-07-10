"""Insert the AIED prior + its fusions into the deck (2026-07-10, Carlos's ask).

AIED = ['MG','1d_DG','BPM','blink_durations'] — the prior pipeline's 4-feature
gaze baseline (features_label="AIED"), same features all splits; distinct from
AIxVR (per-split, per the AIxVR paper). Cells from c2_paper_table.csv.

  Slides 'Modality Ablation — Overall/Dyads/Triads' (12–14): 3 rows inserted
    per slide — AIED after AIxVR, audio+AIED after audio+AIxVR,
    linguistic+AIED after linguistic+AIxVR.
  Slide 'Detector vs Priors' (20): AIED row into the Sets table,
    audio+AIED / ling+AIED into the Fusions table (acc (d) + sig markers).

ROW-INSERT (XML tr copy), NOT rebuild — preserves manual edits (e.g. the
hand-bolded highest-d rows). If patcher 29 is ever re-run (it rebuilds 12–14),
run this one again afterwards. Idempotent: skips slides that already have AIED.

Run from repo root:  python3 discover/presentation_scripts/31_aied_rows.py
"""
import shutil
import os
from copy import deepcopy
import pandas as pd
from pptx import Presentation
from pptx.util import Pt
from pptx.oxml.ns import qn

SRC = "scratch/2026_March_ViLearn_Planning.pptx"
BAK = "scratch/2026_March_ViLearn_Planning_pre-aied_2026-07-10.pptx"
if not os.path.exists(BAK):
    shutil.copyfile(SRC, BAK)
    print(f"backup -> {BAK}")

prs = Presentation(SRC)
c2 = pd.read_csv("scratch/te_adoption/a1_out/c2_paper_table.csv")


def r1(sp, det):
    m = c2[(c2.split == sp) & (c2.detector == det) & (c2["rank"] == 1)]
    if not len(m):
        raise KeyError(f"missing cell {sp}/{det}")
    return m.iloc[0]


def pfmt(p):
    if pd.isna(p) or p == "":
        return "—"
    p = float(p)
    return "<.001" if p < 0.001 else f"{p:.3f}"


def find_slide(prefix):
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip().startswith(prefix):
                return s
    return None


def row_texts_ablation(sp, det, name):
    r = r1(sp, det)
    return [name, str(int(r.n_feat)), r.model, f"{r.acc_mean:.3f} ± {r.acc_std:.3f}",
            f"{r.cohen_d_maj:.2f}", pfmt(r.perm_p), pfmt(r.p_ttest_maj)]


def marker_cell(sp, det):
    r = r1(sp, det)
    m = ("*" if (pd.notna(r.perm_p) and float(r.perm_p) < 0.05 and r.perm_p != "") else "") + \
        ("†" if float(r.p_ttest_maj) < 0.05 else "")
    return f"{r.acc_mean:.3f} ({r.cohen_d_maj:.2f}){m}"


def insert_row_after(table, anchor_idx, texts):
    """Deepcopy the anchor <a:tr>, place it right after, set cell texts keeping
    the anchor row's fonts (anchor rows are never the bolded best rows)."""
    tbl = table._tbl
    src = tbl.tr_lst[anchor_idx]
    new = deepcopy(src)
    src.addnext(new)
    row_idx = anchor_idx + 1
    # template font from the anchor row's first cell run
    tmpl = table.cell(anchor_idx, 0).text_frame.paragraphs[0].runs
    size = tmpl[0].font.size if tmpl and tmpl[0].font.size else Pt(9)
    for j, t in enumerate(texts):
        cell = table.cell(row_idx, j)
        cell.text = str(t)
        for run in cell.text_frame.paragraphs[0].runs:
            run.font.size = size
            run.font.bold = False
    return row_idx


def find_row(table, prefix):
    for i in range(len(table.rows)):
        if table.cell(i, 0).text.strip().replace(" ", "").lower().startswith(prefix.replace(" ", "").lower()):
            return i
    return None


# --- slides 12-14: modality ablation per split -------------------------------
for sp, lab in (("All", "Overall"), ("D", "Dyads"), ("T", "Triads")):
    s = find_slide(f"Modality Ablation — {lab}")
    if s is None:
        print(f"WARN: ablation slide {lab} not found")
        continue
    tbl = next(sh.table for sh in s.shapes if sh.has_table)
    if find_row(tbl, "AIED") is not None:
        print(f"{lab}: AIED rows already present, skip")
        continue
    for anchor, det, name in (("AIxVR", "AIED", "AIED (prior gaze, 4 feat all splits)"),
                              ("audio + AIxVR", "audio+AIED", "audio + AIED"),
                              ("linguistic + AIxVR", "linguistic+AIED", "linguistic + AIED")):
        ai = find_row(tbl, anchor)
        if ai is None:
            print(f"WARN: anchor '{anchor}' not found on {lab}")
            continue
        insert_row_after(tbl, ai, row_texts_ablation(sp, det, name))
    print(f"{lab}: 3 AIED rows inserted")

# --- slide 20: Detector vs Priors (two tables) --------------------------------
s20 = find_slide("Detector vs Priors")
if s20 is None:
    print("WARN: Detector vs Priors slide not found")
else:
    tables = [sh.table for sh in s20.shapes if sh.has_table]
    sets_tbl = next((t for t in tables if find_row(t, "AIxVR") is not None), None)
    fus_tbl = next((t for t in tables if find_row(t, "audio+AIxVR") is not None), None)
    if sets_tbl is not None and find_row(sets_tbl, "AIED") is None:
        ai = find_row(sets_tbl, "AIxVR")
        insert_row_after(sets_tbl, ai,
                         ["AIED", "4", marker_cell("All", "AIED"),
                          marker_cell("D", "AIED"), marker_cell("T", "AIED")])
        print("slide 20 Sets: AIED row inserted")
    else:
        print("slide 20 Sets: skip (present or table not found)")
    if fus_tbl is not None and find_row(fus_tbl, "audio+AIED") is None:
        for anchor, det, name, nf in (("audio+AIxVR", "audio+AIED", "audio+AIED", "7"),
                                      ("ling+AIxVR", "linguistic+AIED", "ling+AIED", "11")):
            ai = find_row(fus_tbl, anchor)
            if ai is None:
                print(f"WARN: fusion anchor '{anchor}' not found")
                continue
            insert_row_after(fus_tbl, ai,
                             [name, nf, marker_cell("All", det),
                              marker_cell("D", det), marker_cell("T", det)])
        print("slide 20 Fusions: AIED fusion rows inserted")
    else:
        print("slide 20 Fusions: skip (present or table not found)")

# --- slide 12 bullet: best-fusion claim now a tie (add-once) ------------------
s12 = find_slide("Modality Ablation — Overall")
OLD = ("Linguistic is the strongest single set (0.771); linguistic+GazexSpeaking the "
       "strongest fusion (0.786). AIxVR is dominated everywhere, incl. its linguistic "
       "fusion (0.741 < 0.771).")
NEW = ("Linguistic is the strongest single set (0.771). Strongest fusions: "
       "linguistic+AIED 0.794 (d=1.81) and linguistic+GazexSpeaking 0.786 (d=1.33) — "
       "statistically tied (paired p=.83): with a text backbone the exact gaze set "
       "barely matters. AIxVR stays dominated everywhere (ling+AIxVR 0.741 < 0.771); "
       "AIED (prior 4-feature gaze set) sits between AIxVR and GazexSpeaking as a single set.")
if s12 is not None:
    done = False
    for sh in s12.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            full = "".join(r.text for r in p.runs)
            if full == OLD:
                for r in p.runs[1:]:
                    r.text = ""
                p.runs[0].text = NEW
                done = True
    print("slide 12 bullet updated" if done else "slide 12 bullet: pattern not found (already updated?)")

prs.save(SRC)
print(f"saved {SRC}")
