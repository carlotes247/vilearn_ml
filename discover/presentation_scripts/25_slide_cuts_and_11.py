"""Deck cuts + slide-11 table rewrite (saveli, 2026-07-07 deck review).

- Slide 11 'Feature Modalities': table rewritten in place to the final paper sets
  (linguistic v2 / audio affect / gaze / fusions+embeddings); the appended
  'Paper linguistic set' sentence is removed (table carries it now); bottom text
  notes that ALL prediction is 60 s group-window level (segment-level dropped).
- DELETE slides: 'What Changed Since Last Week' (12), 'Prior Model — Corrected' (14),
  'Why 60 s Windows' (16) — process slides, out of the paper cut.
  Deletion via drop_rel + sldIdLst removal (no orphaned parts in the zip).

Run from repo root:  python3 discover/presentation_scripts/25_slide_cuts_and_11.py
"""
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Pt

SRC = "scratch/2026_March_ViLearn_Planning.pptx"
prs = Presentation(SRC)


def find_slide(prefix):
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip().startswith(prefix):
                return s
    return None


def set_cell(cell, text, size=Pt(11), bold=False):
    cell.text = str(text)
    for p in cell.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = size
            r.font.bold = bold


# ---------- slide 11 ----------
s11 = find_slide("Feature Modalities")
tbl = next(sh.table for sh in s11.shapes if sh.has_table)
rows = [
    ("Linguistic (7, group-window)",
     "segments/window, words/window, avg word length, question ratio, speech ratio, "
     "mean segment duration, unfinished ratio — raw ASR segments, no sentiment model"),
    ("Audio (affect)", "arousal, dominance, valence (group-level v/a/d, derived from audio)"),
    ("Gaze (prior work)", "GazexSpeaking: BPM, blink durations, gaze on speaker / on silent (same all splits); "
     "AIxVR per-split sets as second prior"),
    ("Fusions / tested & rejected", "concatenations (audio+linguistic, linguistic+gaze, all three); "
     "openSMILE 88, emoW2V + sentiment embeddings all lose to handcrafted (slides on extended sets)"),
]
for i, (a, b) in enumerate(rows, 1):
    set_cell(tbl.cell(i, 0), a, bold=True)
    set_cell(tbl.cell(i, 1), b, size=Pt(10))
print("slide 11: table rewritten")

MARK = "Paper linguistic set"
for sh in s11.shapes:
    if sh.has_text_frame and MARK in sh.text_frame.text:
        tf = sh.text_frame
        # drop the appended paragraph, refresh the tail sentence
        for p in list(tf.paragraphs):
            if MARK in p.text:
                p._p.getparent().remove(p._p)
        base = tf.paragraphs[0]
        base.text = ("Audio is group-shared (one group.audio.wav); the transcript is diarized from it. "
                     "Speaking attributes the shared signal to each role → folded into linguistic. "
                     "v/a/d derived from audio → audio. All prediction is at the 60 s group-window level — "
                     "segment-level and per-role prediction dropped with the individual-TE track.")
        for r in base.runs:
            r.font.size = Pt(11)
        print("slide 11: appended sentence removed, bottom text refreshed")
        break

# ---------- delete slides ----------
for prefix in ("What Changed Since Last Week", "Prior Model — Corrected", "Why 60 s Windows"):
    s = find_slide(prefix)
    if s is None:
        print(f"'{prefix}' not found — already cut?")
        continue
    rId = None
    for k, rel in prs.part.rels.items():
        if rel.reltype.endswith("/slide") and rel.target_part is s.part:
            rId = k
            break
    lst = prs.slides._sldIdLst
    for sldId in list(lst):
        if sldId.get(qn("r:id")) == rId:
            lst.remove(sldId)
    prs.part.drop_rel(rId)
    print(f"deleted '{prefix}'")

prs.save(SRC)
print(f"saved {SRC}: {len(prs.slides)} slides")
