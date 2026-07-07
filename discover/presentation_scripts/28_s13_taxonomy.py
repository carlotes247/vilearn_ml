"""Slide 13: drop single/fused-set taxonomy (saveli: GxS itself fuses gaze x speaking,
AIxVR fuses gaze + blink streams — nothing is cleanly 'single'). Rows relabeled
prior / ours / combined; terminology bullet rewritten.

Run from repo root:  python3 discover/presentation_scripts/28_s13_taxonomy.py
"""
from pptx import Presentation
from pptx.util import Pt

SRC = "scratch/2026_March_ViLearn_Planning.pptx"
prs = Presentation(SRC)


def find_slide(prefix):
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip().startswith(prefix):
                return s
    return None


def set_cell(cell, text, size=Pt(9)):
    cell.text = text
    cell.text_frame.paragraphs[0].runs[0].font.size = size


s13 = find_slide("Main Results — Group-Level")
tbl = next(sh.table for sh in s13.shapes if sh.has_table)
RELABEL = {
    "best single set: linguistic": "ours: linguistic",
    "prior gaze (GxS) = best single set": "prior gaze (GxS) — also best set",
    "best fused set: linguistic+GazexSpeaking": "combined: linguistic+GxS",
    "best fused set: audio+linguistic": "combined: audio+linguistic",
}
for row in tbl.rows:
    lab = row.cells[1].text.strip()
    if lab in RELABEL:
        set_cell(row.cells[1], RELABEL[lab])
print("slide 13 rows relabeled")

OLD_MARK = "single set ≠ single modality"
for sh in s13.shapes:
    if sh.has_text_frame and OLD_MARK in sh.text_frame.text:
        for p in sh.text_frame.paragraphs:
            if OLD_MARK in p.text:
                new = ("Rows are feature SETS, not modalities: GxS itself fuses eye-tracking with "
                       "speaking status (diarization), AIxVR fuses gaze and blink streams — so we "
                       "label sets by name (prior / ours / combined), not by modality count.")
                p.text = new
                for run in p.runs:
                    run.font.size = Pt(10)
        print("terminology bullet rewritten")
        break

prs.save(SRC)
print("saved", SRC)
