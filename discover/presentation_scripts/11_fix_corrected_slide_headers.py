"""Fix the corrected-prior slide: the two non-published columns differ only by
feature scaling, NOT by CV (the nested CV — adopted from the prior — is used in
both). Relabel headers + rewrite the note. Idempotent. Run after 10."""
from pptx import Presentation
from pptx.util import Pt

SRC = "discover/2026_March_ViLearn_Planning.pptx"
TITLE = "Prior Baseline — Corrected for the Label Bug"
prs = Presentation(SRC)


def set_cell(cell, text, sz=12, bold=False):
    cell.text = str(text)
    r = cell.text_frame.paragraphs[0].runs[0]
    r.font.size = Pt(sz); r.font.bold = bold


for s in prs.slides:
    if not any(sh.has_text_frame and sh.text_frame.text.strip() == TITLE for sh in s.shapes):
        continue
    for sh in s.shapes:
        if sh.has_table:
            set_cell(sh.table.cell(0, 2), "Label-fixed · unscaled (prior setup)", bold=True)
            set_cell(sh.table.cell(0, 3), "Label-fixed · scaled (our pipeline)", bold=True)
        elif sh.has_text_frame and sh.text_frame.text.strip().startswith("Same QDA"):
            tf = sh.text_frame
            for p in tf.paragraphs[1:]:
                p._p.getparent().remove(p._p)
            p0 = tf.paragraphs[0]
            for r in list(p0.runs)[1:]:
                r._r.getparent().remove(r._r)
            bullets = [
                "Same QDA, same gaze features, same nested CV (adopted from the prior) in all columns — only the labels and feature scaling change.",
                "Label fix alone (col 2, unscaled = the prior's published SIMPLE setup): 0.60 → 0.67 All. The 60 Hz time-stretch bug depressed published numbers ~0.05–0.08 (13/20 groups; the 7 pure-90 Hz groups replicate exactly).",
                "Cols 2 → 3 differ only by feature scaling (StandardScaler): the prior's published runs were unscaled; our detector pipeline scales. Col 3 is what the results slide compares against.",
                "None reach significance vs chance after Bonferroni (d up to 1.18, but n = 8–20 underpowered).",
            ]
            p0.runs[0].text = bullets[0] if p0.runs else None
            if not p0.runs:
                p0.text = bullets[0]
            p0.runs[0].font.size = Pt(13)
            for text in bullets[1:]:
                p = tf.add_paragraph(); p.text = text
                for r in p.runs:
                    r.font.size = Pt(13)
    break

prs.save(SRC)
print("fixed corrected-prior slide headers + note")
