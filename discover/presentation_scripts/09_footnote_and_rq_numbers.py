"""Results slide: add a methods footnote (fair-comparison construction + window
counts). RQ slide: replace the 'Carlos's...' bullet with a neutral reference to
the depressed slide-2/3 numbers. Idempotent. Run from repo root.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import PP_PLACEHOLDER

SRC = "discover/2026_March_ViLearn_Planning.pptx"
prs = Presentation(SRC)

FOOTNOTE = (
    "Fair comparison: prior gaze features (185 floorlevel windows) re-scored with our "
    "harmonized 2-annotator labels under identical LOSO nested CV; our detector uses our "
    "pipeline windows (197). Counts differ from tail-window trimming; tests paired by group "
    "(n = 20 / 8 / 12)."
)
RQ_OLD_KEY = "Carlos"
RQ_NEW = ("Prior baseline shown on slides 2–3 is depressed by a 60 Hz label time-stretch bug "
          "(e.g. QDA All 0.60 published → 0.67 corrected on harmonized labels)")

# ---- results slide (caption starts with 'Per-fold') ----
for s in prs.slides:
    if any(sh.has_text_frame and sh.text_frame.text.strip().startswith("Per-fold") for sh in s.shapes):
        # idempotent: skip if footnote already present
        if any(sh.has_text_frame and "Fair comparison:" in sh.text_frame.text for sh in s.shapes):
            print("footnote already present")
            break
        tb = s.shapes.add_textbox(Inches(0.5), Inches(6.7), Inches(12.2), Inches(0.8))
        tb.text_frame.word_wrap = True
        p = tb.text_frame.paragraphs[0]
        p.text = FOOTNOTE
        p.runs[0].font.size = Pt(10)
        p.runs[0].font.italic = True
        print("added footnote to results slide")
        break

# ---- RQ slide: replace the Carlos bullet ----
for s in prs.slides:
    body = None
    is_rq = any(sh.has_text_frame and "Answers to Research Questions" in sh.text_frame.text for sh in s.shapes)
    if not is_rq:
        continue
    for ph in s.placeholders:
        if ph.placeholder_format.type != PP_PLACEHOLDER.TITLE:
            body = ph
    if body is None:
        continue
    for p in body.text_frame.paragraphs:
        if RQ_OLD_KEY in p.text:
            sz = p.runs[0].font.size if p.runs else Pt(12)
            for r in list(p.runs)[1:]:
                r._r.getparent().remove(r._r)
            if p.runs:
                p.runs[0].text = RQ_NEW
                p.runs[0].font.size = sz
            else:
                p.text = RQ_NEW
            print("replaced RQ Carlos bullet")
    break

prs.save(SRC)
print("done")
