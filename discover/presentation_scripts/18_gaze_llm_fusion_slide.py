"""Append a gaze x speaking x LLM fusion slide. GazexSpeaking-6 (consistent all splits)
fused with the LLM rubric markers (think pass) and v/a/d. gaze+LLM = best detector on
All (0.86) and Dyads (0.87); triple gaze+vad+LLM best on Triads (0.85). Idempotent."""
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt

SRC = "discover/2026_March_ViLearn_Planning.pptx"
TITLE = "Best Config: Gaze×Speaking × LLM Rubric (fusion)"
prs = Presentation(SRC)
if any(any(sh.has_text_frame and sh.text_frame.text.strip() == TITLE for sh in s.shapes) for s in prs.slides):
    print("slide already present"); raise SystemExit

r = pd.read_csv("scratch/te_adoption/a1_out/gaze_llm_fusion.csv")
piv = r.pivot_table(index="split", columns="features", values="acc", aggfunc="first")

s = prs.slides.add_slide(prs.slide_layouts[7])
s.shapes.title.text = TITLE
header = ["Split", "Gaze (GxS-6)", "LLM", "Gaze+vad", "Gaze+LLM", "Gaze+vad+LLM"]
order = ["gaze", "llm", "gaze+vad", "gaze+llm", "gaze+vad+llm"]
labels = {"all": "All", "D": "Dyads", "T": "Triads"}
rows = [[labels[sp]] + [f"{piv.loc[sp, c]:.2f}" for c in order] for sp in ["all", "D", "T"]]
widths = [1.4, 2.0, 1.2, 1.7, 1.8, 2.2]
tbl = s.shapes.add_table(len(rows) + 1, len(header), Inches(0.5), Inches(1.4),
                         Inches(sum(widths)), Inches(0.45 * (len(rows) + 1))).table
for j, wd in enumerate(widths):
    tbl.columns[j].width = Inches(wd)
for j, h in enumerate(header):
    c = tbl.cell(0, j); c.text = h
    c.text_frame.paragraphs[0].runs[0].font.size = Pt(12); c.text_frame.paragraphs[0].runs[0].font.bold = True
for i, row in enumerate(rows, 1):
    for j, v in enumerate(row):
        c = tbl.cell(i, j); c.text = str(v)
        run = c.text_frame.paragraphs[0].runs[0]; run.font.size = Pt(12)
        run.font.bold = (header[j] == "Gaze+LLM" and row[0] in ("All", "Dyads")) or \
                        (header[j] == "Gaze+vad+LLM" and row[0] == "Triads")

tb = s.shapes.add_textbox(Inches(0.5), Inches(3.5), Inches(12.3), Inches(3.4))
tb.text_frame.word_wrap = True
bullets = [
    "Fuses the prior gaze detector (GazexSpeaking, 6 features, SAME set all splits) with the LLM rubric markers (4 axes, thinking pass) — consistent feature sets, no per-split cherry-picking.",
    "Gaze+LLM = BEST detector: All 0.86, Dyads 0.87 (both perm p≈.003) — beats every prior, every single modality, and the old gaze+vad fusion (0.75/0.81). On Dyads it also turns gaze-alone's ns result (p=.083) significant.",
    "Gaze and rubric-text are COMPLEMENTARY: gaze captures attention / turn-taking, the LLM captures content depth & reasoning — fusing them is additive, not redundant.",
    "v/a/d affect adds nothing once the LLM is in (All/Dyads: gaze+LLM ≥ gaze+vad+LLM). Only Triads benefits from the full triple (0.85 vs 0.83) — where gaze is weakest, audio still helps.",
    "Group-level, nested LOSO + permutation, clean 2-annotator labels. LLM = local Qwen3.6-27B (thinking mode; outperforms non-think on the marker features).",
]
for i, t in enumerate(bullets):
    p = tb.text_frame.paragraphs[0] if i == 0 else tb.text_frame.add_paragraph()
    p.text = t
    for run in p.runs:
        run.font.size = Pt(11)

prs.save(SRC)
print(f"gaze×LLM fusion slide added; {len(prs.slides)} slides")
