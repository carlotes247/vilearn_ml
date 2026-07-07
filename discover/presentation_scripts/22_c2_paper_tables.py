"""C2 deck patch (2026-07-06 meeting): effect sizes on slides + extended pre-LLM sets.

Targets scratch/2026_March_ViLearn_Planning.pptx (make a backup first). Idempotent.

Changes:
  (1) Slide 21 'Detector vs 3 Priors + Fusion': table REBUILT from c2_paper_table.csv —
      adds Acc±SD, Cohen's d vs majority, both p's (perm / t-test). Same sets as before.
  (2) NEW slide inserted after slide 22: linguistic + openSMILE(88) + openSMILE_pca +
      audio+openSMILE under the canonical group-level protocol (Carlos's missing sets).
  (3) Slide 22 (openSMILE vs affect): footnote pointing at the canonical rerun.

Run from repo root:  python3 discover/presentation_scripts/22_c2_paper_tables.py
"""
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt

SRC = "scratch/2026_March_ViLearn_Planning.pptx"
prs = Presentation(SRC)
c2 = pd.read_csv("scratch/te_adoption/a1_out/c2_paper_table.csv")

LAB = {"All": "Overall", "D": "Dyads", "T": "Triads"}


def find_slide(prefix):
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip().startswith(prefix):
                return s
    return None


def slide_index(s):
    return list(prs.slides).index(s)


def set_cell(cell, text, size=Pt(9), bold=False):
    cell.text = str(text)
    r = cell.text_frame.paragraphs[0].runs[0]
    r.font.size = size
    r.font.bold = bold


def pfmt(p):
    if pd.isna(p) or p == "":
        return "—"
    p = float(p)
    return "<.001" if p < 0.001 else f"{p:.3f}"


def r1(sp, det):
    """rank-1 row for (split, detector)."""
    m = c2[(c2.split == sp) & (c2.detector == det) & (c2["rank"] == 1)]
    return m.iloc[0] if len(m) else None


def table_row(sp, det, name=None):
    r = c2[(c2.split == sp) & (c2.detector == det) & (c2["rank"] == 1)].iloc[0]
    if det == "baseline_majority":
        return [LAB[sp], "majority baseline", "—", f"{r.acc_mean:.3f} ± {r.acc_std:.3f}", "—", "—", "—"]
    return [LAB[sp], name or det, r.model, f"{r.acc_mean:.3f} ± {r.acc_std:.3f}",
            f"{r.cohen_d_maj:.2f}", pfmt(r.perm_p), pfmt(r.p_ttest_maj)]


HEADER = ["Split", "Feature set", "Best model", "Acc ± SD", "d", "p perm", "p t-test"]
WIDTHS = [1.0, 2.3, 1.2, 1.9, 0.8, 1.0, 1.0]


def build_table(s, rows, top=1.15, font=Pt(9)):
    tbl = s.shapes.add_table(len(rows) + 1, len(HEADER), Inches(0.4), Inches(top),
                             Inches(sum(WIDTHS)), Inches(0.31 * (len(rows) + 1))).table
    for j, wd in enumerate(WIDTHS):
        tbl.columns[j].width = Inches(wd)
    for j, h in enumerate(HEADER):
        set_cell(tbl.cell(0, j), h, size=Pt(10), bold=True)
    for i, row in enumerate(rows, 1):
        best_of_split = row[1] != "majority baseline" and row == max(
            [r for r in rows if r[0] == row[0] and r[1] != "majority baseline"],
            key=lambda r: float(r[3].split(" ±")[0]))
        for j, v in enumerate(row):
            set_cell(tbl.cell(i, j), v, size=font, bold=best_of_split and j in (1, 3))
    return tbl


# ---------- (1) slide 21 rebuild with effect sizes ----------
ORIG = ["AIxVR", "GazexSpeaking", "audio", "audio+AIxVR", "audio+GazexSpeaking"]
s21 = find_slide("Detector vs 3 Priors")
if s21:
    title_sh = s21.shapes[0]
    for sh in list(s21.shapes)[1:]:
        sh._element.getparent().remove(sh._element)
    title_sh.text_frame.text = "Detector vs 3 Priors + Fusion (effect sizes, canonical protocol)"
    rows = []
    for sp in ("All", "D", "T"):
        rows.append(table_row(sp, "baseline_majority"))
        rows += [table_row(sp, det) for det in ORIG]
    build_table(s21, rows)
    nb = s21.shapes.add_textbox(Inches(0.4), Inches(6.9), Inches(12.5), Inches(0.5))
    nb.text_frame.word_wrap = True
    nb.text_frame.paragraphs[0].text = (
        "Group-level, 185 windows, clean 2-annotator labels, nested LOSO. d = Cohen's d vs majority floor "
        "(per-fold accs); p perm = 300 label permutations (best model, default hyperparams); p t-test = "
        "one-sample t vs majority. Per-fold scores: c2_foldscores.csv. Second-best model per set in c2_paper_table.csv.")
    nb.text_frame.paragraphs[0].runs[0].font.size = Pt(8)
    nb.text_frame.paragraphs[0].runs[0].font.italic = True
    print("slide 21 rebuilt with ±SD, d, both p")
else:
    print("slide 21 NOT found — skipped")

# ---------- (2) new extended-sets slide after slide 22 ----------
NEW_TITLE = "Extended Feature Sets — Linguistic & Raw Acoustics (canonical protocol)"
EXT = [("linguistic", "linguistic (7 transcript stats)"),
       ("openSMILE", "openSMILE (88 eGeMAPS)"),
       ("openSMILE_pca", "openSMILE + PCA (tuned 5/10/20)"),
       ("audio+openSMILE", "audio(v/a/d)+openSMILE")]
if not find_slide(NEW_TITLE):
    s = prs.slides.add_slide(prs.slide_layouts[7])
    s.shapes.title.text = NEW_TITLE
    rows = []
    for sp in ("All", "D", "T"):
        rows.append(table_row(sp, "baseline_majority"))
        for det, name in EXT:
            if r1(sp, det) is not None:
                rows.append(table_row(sp, det, name))
    build_table(s, rows)
    lin_all = r1("All", "linguistic")
    bullets = [
        f"NEW: general-linguistic set (word count/rate/length, question rate, speaking seconds, segment stats — "
        f"no model-derived sentiment) on the SAME 185-window protocol as gaze/audio: "
        f"Overall {lin_all.acc_mean:.3f} ± {lin_all.acc_std:.3f} (d={lin_all.cohen_d_maj:.2f}) — "
        "strongest single pre-LLM modality Overall.",
        "openSMILE raw acoustics (88 eGeMAPS functionals) fail at group level even with PCA tuned inside the "
        "nested pipeline (no leakage) — confirms the derive-grid finding (slide 22): derived affect (v/a/d) "
        "beats raw acoustics.",
        "QDA is excluded where the class covariance is singular (88 features > per-class n); remaining panel "
        "(SVM / LogReg / NB) unaffected.",
    ]
    tb = s.shapes.add_textbox(Inches(0.4), Inches(1.15 + 0.31 * (len(rows) + 1) + 0.25),
                              Inches(12.5), Inches(2.2))
    tb.text_frame.word_wrap = True
    for i, t in enumerate(bullets):
        p = tb.text_frame.paragraphs[0] if i == 0 else tb.text_frame.add_paragraph()
        p.text = t
        for run in p.runs:
            run.font.size = Pt(10)
    # move from end to right after slide 22 (index 22)
    lst = prs.slides._sldIdLst
    el = list(lst)[-1]
    lst.remove(el)
    lst.insert(22, el)
    print(f"inserted '{NEW_TITLE}' as slide 23")
else:
    print("extended-sets slide already present — skip")

# ---------- (3) slide 22 footnote ----------
s22 = find_slide("Raw Acoustics (openSMILE) vs Affect")
if s22:
    NOTE = "Canonical group-level rerun (185 windows, next slide): openSMILE still loses"
    if not any(sh.has_text_frame and NOTE in sh.text_frame.text for sh in s22.shapes):
        os_pca = {sp: r1(sp, "openSMILE_pca") for sp in ("All", "D", "T")}
        nb = s22.shapes.add_textbox(Inches(0.4), Inches(6.9), Inches(12.5), Inches(0.5))
        nb.text_frame.word_wrap = True
        nb.text_frame.paragraphs[0].text = (
            f"{NOTE} — best (PCA-tuned) {os_pca['All'].acc_mean:.3f}/{os_pca['D'].acc_mean:.3f}/"
            f"{os_pca['T'].acc_mean:.3f} (All/D/T) vs audio v/a/d 0.663/0.603/0.714.")
        nb.text_frame.paragraphs[0].runs[0].font.size = Pt(9)
        nb.text_frame.paragraphs[0].runs[0].font.italic = True
        print("slide 22 footnote added")
    else:
        print("slide 22 footnote already present — skip")
else:
    print("slide 22 NOT found — skipped")

prs.save(SRC)
print(f"\nsaved {SRC}: {len(prs.slides)} slides")
