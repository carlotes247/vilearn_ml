"""Append a '3 priors + audio + fusion' comparison slide (group-level, clean labels)
showing BOTH significance tests (permutation + majority one-sample t-test) since the
method is undecided. Idempotent."""
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt

SRC = "discover/2026_March_ViLearn_Planning.pptx"
TITLE = "Our Detector vs 3 Priors + Fusion (group-level, clean labels)"
prs = Presentation(SRC)
if any(any(sh.has_text_frame and sh.text_frame.text.strip() == TITLE for sh in s.shapes) for s in prs.slides):
    print("slide already present"); raise SystemExit

t = pd.read_csv("scratch/te_adoption/a1_out/comprehensive_table.csv")


def mark(p):
    if pd.isna(p) or p == "-":
        return "—"
    p = float(p)
    return f"{p:.3f}{'✓' if p < 0.05 else '✗'}"


s = prs.slides.add_slide(prs.slide_layouts[7])
s.shapes.title.text = TITLE
header = ["Split", "Detector", "acc", "maj", "perm p", "t-test p (vs maj)"]
widths = [1.3, 3.4, 1.1, 1.1, 2.0, 2.6]
rows = []
for _, r in t.iterrows():
    rows.append([r.split, r.detector, "—" if r.detector.startswith("dummy") else f"{r.acc:.2f}",
                 f"{r.majority:.2f}", mark(r.perm_p), mark(r.get("p_ttest_majority", "-"))])
tbl = s.shapes.add_table(len(rows) + 1, len(header), Inches(0.3), Inches(1.1),
                         Inches(sum(widths)), Inches(0.3 * (len(rows) + 1))).table
for j, wd in enumerate(widths):
    tbl.columns[j].width = Inches(wd)
for j, h in enumerate(header):
    c = tbl.cell(0, j); c.text = h
    c.text_frame.paragraphs[0].runs[0].font.size = Pt(10); c.text_frame.paragraphs[0].runs[0].font.bold = True
for i, row in enumerate(rows, 1):
    for j, v in enumerate(row):
        c = tbl.cell(i, j); c.text = str(v)
        c.text_frame.paragraphs[0].runs[0].font.size = Pt(9)

tb = s.shapes.add_textbox(Inches(0.3), Inches(6.5), Inches(12.6), Inches(0.9))
tb.text_frame.word_wrap = True
p = tb.text_frame.paragraphs[0]
p.text = ("3 priors: dummy(majority) / AIxVR (gaze, per-split feats) / GazexSpeaking (gaze×speaking). "
          "audio = our gaze-free v/a/d. Two tests shown (method UNDECIDED): permutation null = model on "
          "shuffled labels (lenient); t-test null = majority floor (stricter) → they disagree (esp. Triads & "
          "Dyads). Highlights: fusion audio+GazexSpeaking best on All (0.75) & Dyads (0.81); audio alone wins "
          "Triads (0.71, gaze-free). Dyads n=8 — unstable, caveat.")
p.runs[0].font.size = Pt(10); p.runs[0].font.italic = True

prs.save(SRC)
print(f"3-prior slide added; {len(prs.slides)} slides")
