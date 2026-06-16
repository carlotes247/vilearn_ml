"""Slide 2 redone PROPERLY: prior gaze models on CLEAN labels, group-level (185),
scaled nested CV. Fills all cells, clean majority baseline. Significance is now
method-dependent (uniform/majority/permutation differ) -> footnote points to the
Methodology Review slides instead of hard red boxes. Idempotent."""
from pptx import Presentation
from pptx.util import Inches, Pt

SRC = "discover/2026_March_ViLearn_Planning.pptx"
prs = Presentation(SRC)
s = prs.slides[1]

# clean group-level numbers (scaled nested CV); SD shown as em-dash (per-fold SD large, see Methodology)
# col map: 1 D-M, 2 D-SD, 3 T-M, 4 T-SD, 5 All-M, 6 All-SD
ROWS = {
    "Majority baseline": ["0.60", "—", "0.54", "—", "0.52", "—"],
    "SVM": ["0.83", "—", "0.56", "—", "0.73", "—"],
    "Naive Bayes": ["0.72", "—", "0.70", "—", "0.71", "—"],
    "QDA": ["0.77", "—", "0.65", "—", "0.73", "—"],
}


def set_cell(cell, text):
    old = cell.text_frame.paragraphs[0].runs[0] if cell.text_frame.paragraphs[0].runs else None
    sz = old.font.size if old and old.font.size else Pt(12)
    bold = old.font.bold if old else False
    cell.text = str(text)
    r = cell.text_frame.paragraphs[0].runs[0]
    r.font.size = sz; r.font.bold = bool(bold)


for sh in s.shapes:
    if sh.has_table and any("GazexSpeaking" in c.text for c in sh.table.rows[0].cells):
        t = sh.table
        for ri in range(len(t.rows)):
            label = t.cell(ri, 0).text.strip()
            key = "Majority baseline" if label.lower().startswith(("baseline", "majority")) else label
            if key in ROWS:
                if key == "Majority baseline":
                    set_cell(t.cell(ri, 0), "Majority baseline")
                for cj, v in enumerate(ROWS[key], start=1):
                    set_cell(t.cell(ri, cj), v)

# replace footnote
for sh in s.shapes:
    if sh.has_text_frame and "Majority baseline =" in sh.text_frame.text:
        tf = sh.text_frame
        for p in tf.paragraphs[1:]:
            p._p.getparent().remove(p._p)
        p0 = tf.paragraphs[0]
        for r in list(p0.runs)[1:]:
            r._r.getparent().remove(r._r)
        txt = ("Redone on CLEAN 2-annotator labels, group-level (185 windows), scaled nested CV "
               "(was bugged labels + uniform/seed baseline). Majority baseline = deterministic "
               "always-predict-majority. Significance is now METHOD-DEPENDENT (uniform vs majority vs "
               "permutation give different verdicts) — see 'Methodology Review' slides; no fixed red boxes.")
        if p0.runs:
            p0.runs[0].text = txt
        else:
            p0.text = txt
        break

prs.save(SRC)
print("slide 2 redone: clean group-level prior models + method-dependent-sig footnote")
