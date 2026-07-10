"""openSMILE feature-selection results (2026-07-10 experiments d2/d4 + openSMILE_sel).

  Slide 'Extended Feature Sets' (22 after the slide-12 split): table REBUILT with
    the openSMILE_sel row per split + refreshed bullets (also fixes the stale
    pre-v2 linguistic number 0.788 -> 0.771).
  Slide 'Raw Acoustics' (21): canonical-rerun footnote extended with sel numbers.

Findings encoded here:
  - in-CV SelectKBest (log grid 1..88) picks k~4 (All 15/20 folds, T 9/12);
    converges on VARIABILITY functionals (mfcc1V/spectralFlux/loudness stddevNorm,
    F0/loudness range) — acoustic cousin of word_count. Dyads unstable (k 4-32).
  - acc 0.602/0.692/0.602 — selection does not rescue raw acoustics.
  - PCA n_components probe (d2): All picks 5 = grid floor in 20/20 folds.

Targets scratch/2026_March_ViLearn_Planning.pptx. Idempotent: table rebuilt
each run; footnote append is add-once.

Run from repo root:  python3 discover/presentation_scripts/30_opensmile_selection.py
"""
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt

SRC = "scratch/2026_March_ViLearn_Planning.pptx"
prs = Presentation(SRC)
c2 = pd.read_csv("scratch/te_adoption/a1_out/c2_paper_table.csv")
LAB = {"All": "Overall", "D": "Dyads", "T": "Triads"}

ROWS = [
    ("baseline_majority", "majority baseline"),
    ("linguistic", "linguistic (7 segment stats)"),
    ("openSMILE", "openSMILE (88 eGeMAPS)"),
    ("openSMILE_pca", "openSMILE + PCA (tuned 5/10/20)"),
    ("openSMILE_sel", "openSMILE + SelectKBest (k tuned 1–88)"),
    ("audio+openSMILE", "audio(v/a/d)+openSMILE"),
]
HEADER = ["Split", "Feature set", "Best model", "Acc ± SD", "d", "p perm", "p t-test"]
WIDTHS = [1.0, 3.2, 1.2, 1.9, 0.8, 1.0, 1.0]


def find_slide(prefix):
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip().startswith(prefix):
                return s
    return None


def set_cell(cell, text, size=Pt(9), bold=False):
    cell.text = str(text)
    r = cell.text_frame.paragraphs[0].runs[0]
    r.font.size = size
    r.font.bold = bold


def pfmt(p):
    if pd.isna(p) or p == "":
        return "—"
    p = float(p)
    return "<.001" if p < 0.001 else f"{p:.3f}"


def row_for(sp, det, name):
    r = c2[(c2.split == sp) & (c2.detector == det) & (c2["rank"] == 1)].iloc[0]
    if det == "baseline_majority":
        return [LAB[sp], name, "—", f"{r.acc_mean:.3f} ± {r.acc_std:.3f}", "—", "—", "—"]
    return [LAB[sp], name, r.model, f"{r.acc_mean:.3f} ± {r.acc_std:.3f}",
            f"{r.cohen_d_maj:.2f}", pfmt(r.perm_p), pfmt(r.p_ttest_maj)]


BULLETS = [
    "Linguistic set (7 surface transcript stats, v2) on the SAME 185-window protocol: Overall "
    "0.771 ± 0.210 (d=1.20, both tests sig) — strongest single pre-LLM modality Overall.",
    "openSMILE raw acoustics fail at group level under every reduction tried INSIDE the nested "
    "pipeline (no leakage): PCA (tuned 5/10/20) 0.598/0.668/0.540, SelectKBest (k tuned on a "
    "log grid 1–88) 0.602/0.692/0.602 — derived affect (v/a/d) and text stay ahead.",
    "What little acoustic signal exists is VARIABILITY, not level: selection picks k≈4 in most "
    "folds (Overall 15/20, Triads 9/12) and converges on mfcc1V/spectralFlux/loudness stddevNorm "
    "+ F0/loudness range — 'lively speech dynamics', the acoustic cousin of word_count. "
    "Dyads unstable (k 4–32, n=8). PCA probe agrees: Overall picks the 5-component grid floor 20/20.",
    "Full ranking: d4_opensmile_ranking.csv; per-fold picks: c2_selection.csv. QDA excluded on "
    "high-dim sets (singular covariance).",
]
NOTE = ("Group-level, 185 floorlevel windows, clean 2-annotator labels, nested LOSO; "
        "d = Cohen's d vs majority floor; p perm = 300 permutations; "
        "p t-test = one-sample t vs majority. Source: c2_paper_table.csv.")

s = find_slide("Extended Feature Sets")
if s is None:
    raise SystemExit("Extended Feature Sets slide not found")
title_sh = next(sh for sh in s.shapes if sh.has_text_frame)
keep = title_sh._element
for sh in list(s.shapes):
    if sh._element is not keep:
        sh._element.getparent().remove(sh._element)
rows = [row_for(sp, det, name) for sp in ("All", "D", "T") for det, name in ROWS]
tbl = s.shapes.add_table(len(rows) + 1, len(HEADER), Inches(0.4), Inches(1.1),
                         Inches(sum(WIDTHS)), Inches(0.26 * (len(rows) + 1))).table
for j, wd in enumerate(WIDTHS):
    tbl.columns[j].width = Inches(wd)
for j, h in enumerate(HEADER):
    set_cell(tbl.cell(0, j), h, size=Pt(10), bold=True)
for i, row in enumerate(rows, 1):
    for j, v in enumerate(row):
        set_cell(tbl.cell(i, j), v, size=Pt(8), bold=(j == 1 and "SelectKBest" in row[1]))
y = 1.1 + 0.26 * (len(rows) + 1) + 0.2
tb = s.shapes.add_textbox(Inches(0.4), Inches(y), Inches(12.5), Inches(6.7 - y))
tb.text_frame.word_wrap = True
for i, t in enumerate(BULLETS):
    p = tb.text_frame.paragraphs[0] if i == 0 else tb.text_frame.add_paragraph()
    p.text = t
    for run in p.runs:
        run.font.size = Pt(9)
nb = s.shapes.add_textbox(Inches(0.4), Inches(6.9), Inches(12.5), Inches(0.5))
nb.text_frame.word_wrap = True
nb.text_frame.paragraphs[0].text = NOTE
nb.text_frame.paragraphs[0].runs[0].font.size = Pt(8)
nb.text_frame.paragraphs[0].runs[0].font.italic = True
print("rebuilt 'Extended Feature Sets' (18 data rows incl. openSMILE_sel)")

# --- Raw Acoustics footnote: append sel numbers (add-once) ------------------
s19 = find_slide("Raw Acoustics")
MARK = "SelectKBest"
if s19 is not None:
    hit = False
    for sh in s19.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text
        if "Canonical group-level rerun" in t and MARK not in t:
            p = sh.text_frame.paragraphs[-1]
            run = p.runs[-1] if p.runs else p.add_run()
            run.text += (" In-CV SelectKBest (k tuned 1–88, log grid): 0.602/0.692/0.602 — "
                         "selection doesn't rescue raw acoustics either; picks converge on "
                         "loudness/spectral variability (see Extended Feature Sets slide).")
            hit = True
    print("Raw Acoustics footnote extended" if hit else "Raw Acoustics footnote already extended / not found")

prs.save(SRC)
print(f"saved {SRC}")
