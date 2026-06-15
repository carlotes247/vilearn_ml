"""Add a 'What Changed Since Last Week' opener slide, placed right before the
Group-TE results slide. Pre-empts the "why did the numbers drop / lose
significance" question. Idempotent: skips if the slide already exists.
Run from repo root.
"""
from pptx import Presentation
from pptx.util import Inches, Pt

SRC = "discover/2026_March_ViLearn_Planning.pptx"
TITLE = "What Changed Since Last Week"
prs = Presentation(SRC)

# idempotent guard
for s in prs.slides:
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip() == TITLE:
            print("slide already present — nothing to do")
            raise SystemExit

slide = prs.slides.add_slide(prs.slide_layouts[7])  # Title Only
slide.shapes.title.text = TITLE

# --- comparison table ---
header = ["", "Last week", "Now (harmonized)"]
rows = [
    ["Labels", "helen only · 74% high (skew)", "2-annotator mean · ~49% (balanced)"],
    ["Group TE — All", "0.82", "0.68"],
    ["Group TE — Triads", "0.88", "0.70"],
    ["Group TE — Dyads", "0.74", "0.65"],
    ["vs prior gaze QDA", "beats ✓ (d=0.92–1.35)", "n.s. — matches prior (0.73)"],
]
nrows, ncols = len(rows) + 1, len(header)
widths = [2.6, 3.7, 4.2]
tbl = slide.shapes.add_table(nrows, ncols, Inches(0.5), Inches(1.5),
                             Inches(sum(widths)), Inches(0.4 * nrows)).table
for j, w in enumerate(widths):
    tbl.columns[j].width = Inches(w)
for j, h in enumerate(header):
    c = tbl.cell(0, j); c.text = h
    r = c.text_frame.paragraphs[0].runs[0] if c.text_frame.paragraphs[0].runs else None
    if r:
        r.font.size = Pt(13); r.font.bold = True
for i, row in enumerate(rows, 1):
    for j, val in enumerate(row):
        c = tbl.cell(i, j); c.text = str(val)
        run = c.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(12)
        if j == 0:
            run.font.bold = True

# --- explanation bullets ---
tb = slide.shapes.add_textbox(Inches(0.5), Inches(4.5), Inches(11.5), Inches(2.6))
tf = tb.text_frame; tf.word_wrap = True
bullets = [
    "Two artifacts removed — they are independent:",
    "(1) Our labels were single-annotator + 74% high → majority baseline 0.74 inflated last week's accuracy. Adopting the published 2-annotator mean balances to ~49% → honest 0.68.",
    "(2) The prior baseline itself was computed on time-stretch-bugged labels (60 Hz averaging error, 13/20 groups). Re-ran it on clean labels under identical nested CV.",
    "Net: 0.68, statistically tied with the prior (neither wins); double dissociation + gaze-free parity hold. Better surfaced now than at review.",
]
for i, text in enumerate(bullets):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.text = text
    p.level = 0 if i == 0 else 1
    for r in p.runs:
        r.font.size = Pt(14 if i == 0 else 12)
        if i == 0:
            r.font.bold = True

# --- move new slide to just before the results slide (after Modality Ablation) ---
def title_of(s):
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            return sh.text_frame.text.strip()
    return ""


target = None
for i, s in enumerate(prs.slides):
    if "Results + Significance" in " ".join(sh.text_frame.text for sh in s.shapes
                                            if sh.has_text_frame):
        target = i
        break
xml = prs.slides._sldIdLst
ids = list(xml)
new_el = ids[-1]  # appended last
xml.remove(new_el)
xml.insert(target if target is not None else len(ids) - 1, new_el)

prs.save(SRC)
print(f"added '{TITLE}' at position {target} (before results);", len(prs.slides), "slides")
