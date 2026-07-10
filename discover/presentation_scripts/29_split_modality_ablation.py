"""Split the Modality Ablation slide (12) into 3 per-split slides — Overall /
Dyads / Triads — with ALL base sets + fusions from c2_paper_table.csv
(2026-07-10 decision: Tobi wants AIxVR / ling+AIxVR / ling+GxS rows with stats
and effect sizes; one slide per split gives the table room).

  Slide 12  -> rebuilt as "Modality Ablation — Overall"  (11 rows)
  NEW 13    -> "Modality Ablation — Dyads"
  NEW 14    -> "Modality Ablation — Triads"
  Cross-refs shifted +2 where they point past old slide 12:
    old 12 "slides 18/20" -> handled by rebuild (new bullets use new numbers)
    old 13 (now 15) "slides 18 & 20" -> "slides 20 & 22"
    old 20 (now 22) "(slide 19)" -> "(slide 21)"

Targets scratch/2026_March_ViLearn_Planning.pptx (backup _pre-split12_2026-07-10).
Idempotent: per-split slides are rebuilt in place when already present;
cross-ref replacements are exact-string, skip when already applied.

Run from repo root:  python3 discover/presentation_scripts/29_split_modality_ablation.py
"""
import shutil
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt

SRC = "scratch/2026_March_ViLearn_Planning.pptx"
BAK = "scratch/2026_March_ViLearn_Planning_pre-split12_2026-07-10.pptx"
import os
if not os.path.exists(BAK):
    shutil.copyfile(SRC, BAK)
    print(f"backup -> {BAK}")

prs = Presentation(SRC)
c2 = pd.read_csv("scratch/te_adoption/a1_out/c2_paper_table.csv")

SETS = [
    ("baseline_majority", "majority baseline"),
    ("AIxVR", "AIxVR (prior gaze)"),
    ("GazexSpeaking", "GazexSpeaking (prior gaze)"),
    ("audio", "audio (affect v/a/d)"),
    ("linguistic", "linguistic (transcript stats)"),
    ("audio+AIxVR", "audio + AIxVR"),
    ("audio+GazexSpeaking", "audio + GazexSpeaking"),
    ("audio+linguistic", "audio + linguistic"),
    ("linguistic+AIxVR", "linguistic + AIxVR"),
    ("linguistic+GazexSpeaking", "linguistic + GazexSpeaking"),
    ("audio+linguistic+GazexSpeaking", "audio + linguistic + GazexSpeaking"),
]
HEADER = ["Feature set", "#f", "Best model", "Acc ± SD", "d", "p perm", "p t-test"]
WIDTHS = [3.0, 0.7, 1.2, 1.9, 0.8, 1.0, 1.0]
PROTOCOL_NOTE = ("Group-level, 185 floorlevel windows, clean 2-annotator labels, nested LOSO; "
                 "d = Cohen's d vs majority floor; p perm = 300 permutations; "
                 "p t-test = one-sample t vs majority. Source: c2_paper_table.csv.")

SLIDES = {
    "All": ("Modality Ablation — Overall (60 s group windows, nested LOSO)", [
        "Linguistic is the strongest single set (0.771); linguistic+GazexSpeaking the strongest fusion (0.786). AIxVR is dominated everywhere, incl. its linguistic fusion (0.741 < 0.771).",
        "UNIT FIX: earlier version broadcast the group label onto per-role rows — that diluted text features (a silent listener contributed a zero-word row to an engaged group). Every modality is now aggregated over the whole group window: same unit as the target construct.",
        "Individual-TE prediction track ARCHIVED (labels are single-annotator; group TE has the clean 2-annotator mean). This paper = group TE only.",
        "Per-split tables: next two slides. Extended sets (openSMILE, embeddings): slides 20/22.",
    ]),
    "D": ("Modality Ablation — Dyads (60 s group windows, nested LOSO)", [
        "Gaze rules dyads: GazexSpeaking 0.830 best single (t-test sig, perm ns at n=8); linguistic+GazexSpeaking 0.842 is the best dyad detector and the ONLY one significant under both tests.",
        "Audio is flat (0.603, ns both) and drags fusions down (audio+linguistic 0.639 < linguistic 0.680).",
        "n=8 dyad groups — wide variance; treat single-marker cells as trends.",
    ]),
    "T": ("Modality Ablation — Triads (60 s group windows, nested LOSO)", [
        "Gaze-free wins triads: audio+linguistic 0.784 best; linguistic 0.766 best single set. GazexSpeaking drops to 0.695 — gaze needs an unambiguous target; diffuse in 3-person groups.",
        "AIxVR reduces to 1 feature (BPM) in triads — 0.581, uninformative.",
        "Adding gaze to the best fusion hurts: triple 0.757 < audio+linguistic 0.784.",
    ]),
}
ORDER = ["All", "D", "T"]


def find_slide(prefix):
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip().startswith(prefix):
                return s
    return None


def set_cell(cell, text, size=Pt(9), bold=False):
    cell.text = str(text)
    r = cell.text_frame.paragraphs[0].runs[0]
    r.font.size = size
    r.font.bold = bold


def pfmt(p):
    if pd.isna(p) or p == "":
        return "—"
    p = float(p)
    return "<.001" if p < 0.001 else f"{p:.3f}"


def r1(sp, det):
    m = c2[(c2.split == sp) & (c2.detector == det) & (c2["rank"] == 1)]
    return m.iloc[0] if len(m) else None


def row_for(sp, det, name):
    r = r1(sp, det)
    if r is None:
        raise KeyError(f"missing cell {sp}/{det}")
    if det == "baseline_majority":
        return [name, "—", "—", f"{r.acc_mean:.3f} ± {r.acc_std:.3f}", "—", "—", "—"]
    return [name, str(int(r.n_feat)), r.model, f"{r.acc_mean:.3f} ± {r.acc_std:.3f}",
            f"{r.cohen_d_maj:.2f}", pfmt(r.perm_p), pfmt(r.p_ttest_maj)]


def rebuild(slide, title, rows, bullets, note):
    title_sh = next((sh for sh in slide.shapes if sh.has_text_frame), None)
    keep = title_sh._element if title_sh is not None else None
    for sh in list(slide.shapes):
        if sh._element is not keep:
            sh._element.getparent().remove(sh._element)
    if title_sh is None:
        title_sh = slide.shapes.add_textbox(Inches(0.4), Inches(0.3), Inches(12.5), Inches(0.7))
    title_sh.text_frame.text = title
    for run in title_sh.text_frame.paragraphs[0].runs:
        run.font.size = Pt(24)
        run.font.bold = True
    tbl = slide.shapes.add_table(len(rows) + 1, len(HEADER), Inches(0.4), Inches(1.1),
                                 Inches(sum(WIDTHS)), Inches(0.3 * (len(rows) + 1))).table
    for j, wd in enumerate(WIDTHS):
        tbl.columns[j].width = Inches(wd)
    for j, h in enumerate(HEADER):
        set_cell(tbl.cell(0, j), h, size=Pt(10), bold=True)
    real = [r for r in rows if r[2] != "—"]
    best = max(real, key=lambda r: float(r[3].split(" ±")[0]))
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row):
            set_cell(tbl.cell(i, j), v, size=Pt(9), bold=(row is best and j in (0, 3)))
    y = 1.1 + 0.3 * (len(rows) + 1) + 0.25
    tb = slide.shapes.add_textbox(Inches(0.4), Inches(y), Inches(12.5), Inches(6.7 - y))
    tb.text_frame.word_wrap = True
    for i, t in enumerate(bullets):
        p = tb.text_frame.paragraphs[0] if i == 0 else tb.text_frame.add_paragraph()
        p.text = t
        for run in p.runs:
            run.font.size = Pt(10)
    nb = slide.shapes.add_textbox(Inches(0.4), Inches(6.9), Inches(12.5), Inches(0.5))
    nb.text_frame.word_wrap = True
    nb.text_frame.paragraphs[0].text = note
    nb.text_frame.paragraphs[0].runs[0].font.size = Pt(8)
    nb.text_frame.paragraphs[0].runs[0].font.italic = True


# --- slide 12 (Overall): rebuild in place -----------------------------------
base = find_slide("Modality Ablation — Overall") or find_slide("Modality Ablation")
if base is None:
    raise SystemExit("Modality Ablation slide not found")
lst = prs.slides._sldIdLst
base_pos = next(i for i in range(len(list(lst))) if prs.slides[i] is base)

rebuild(base, SLIDES["All"][0], [row_for("All", d, n) for d, n in SETS],
        SLIDES["All"][1], PROTOCOL_NOTE)
print(f"rebuilt slide {base_pos + 1}: Overall")

# --- Dyads / Triads: rebuild if present, else insert after base -------------
insert_at = base_pos + 1
for sp in ("D", "T"):
    title, bullets = SLIDES[sp]
    prefix = title.split(" (")[0]
    s = find_slide(prefix)
    created = False
    if s is None:
        s = prs.slides.add_slide(prs.slide_layouts[7])
        created = True
    rebuild(s, title, [row_for(sp, d, n) for d, n in SETS], bullets, PROTOCOL_NOTE)
    if created:
        el = list(lst)[-1]
        lst.remove(el)
        lst.insert(insert_at, el)
        print(f"inserted '{prefix}' as slide {insert_at + 1}")
    insert_at += 1

# --- cross-ref shifts (+2 for refs past old slide 12), add-once -------------
FIXES = [
    ("Main Results — Group-Level", "slides 18 & 20", "slides 20 & 22"),
    ("Extended Feature Sets", "(slide 19)", "(slide 21)"),
]
for prefix, old, new in FIXES:
    s = find_slide(prefix)
    if s is None:
        print(f"WARN: slide '{prefix}' not found for cross-ref fix")
        continue
    done = False
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if old in r.text:
                    r.text = r.text.replace(old, new)
                    done = True
    print(f"cross-ref '{prefix}': {old} -> {new}" if done else
          f"cross-ref '{prefix}': '{old}' not present (already fixed?)")

prs.save(SRC)
print(f"saved {SRC}: {len(list(prs.slides._sldIdLst))} slides")
