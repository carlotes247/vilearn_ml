"""Slide 18 rebuild: two compact side-by-side tables (saveli, 2026-07-07).

Left  = single feature sets (baseline, AIxVR, GxS, audio, linguistic)
Right = all fusions (audio+AIxVR, audio+GxS, audio+ling, ling+AIxVR, ling+GxS, triple)
Cells = accuracy (3dp) per split; BOLD = significant under BOTH tests (perm & t<.05).
Full stats (±SD, d, both p) stay in c2_paper_table.csv + slides 12/13/20.

Run from repo root:  python3 discover/presentation_scripts/26_slide18_two_tables.py
"""
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt

SRC = "scratch/2026_March_ViLearn_Planning.pptx"
prs = Presentation(SRC)
c2 = pd.read_csv("scratch/te_adoption/a1_out/c2_paper_table.csv")

SINGLES = [("baseline_majority", "majority baseline"), ("AIxVR", "AIxVR gaze"),
           ("GazexSpeaking", "GazexSpeaking gaze"), ("audio", "audio (v/a/d)"),
           ("linguistic", "linguistic (7)")]
FUSIONS = [("audio+AIxVR", "audio+AIxVR"), ("audio+GazexSpeaking", "audio+GxS"),
           ("audio+linguistic", "audio+linguistic"), ("linguistic+AIxVR", "ling+AIxVR"),
           ("linguistic+GazexSpeaking", "ling+GxS"),
           ("audio+linguistic+GazexSpeaking", "audio+ling+GxS")]


def r1(sp, det):
    m = c2[(c2.split == sp) & (c2.detector == det) & (c2["rank"] == 1)]
    return m.iloc[0] if len(m) else None


def find_slide(prefix):
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip().startswith(prefix):
                return s
    return None


def set_cell(cell, text, size=Pt(11), bold=False):
    cell.text = str(text)
    r = cell.text_frame.paragraphs[0].runs[0]
    r.font.size = size
    r.font.bold = bold


def add_matrix(slide, x, title, sets):
    tb = slide.shapes.add_textbox(Inches(x), Inches(1.05), Inches(6.0), Inches(0.4))
    tb.text_frame.paragraphs[0].text = title
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    tb.text_frame.paragraphs[0].runs[0].font.bold = True
    widths = [2.45, 1.05, 1.05, 1.05]
    tbl = slide.shapes.add_table(len(sets) + 1, 4, Inches(x), Inches(1.5),
                                 Inches(sum(widths)), Inches(0.34 * (len(sets) + 1))).table
    for j, wd in enumerate(widths):
        tbl.columns[j].width = Inches(wd)
    for j, h in enumerate(["Feature set", "All", "Dyads", "Triads"]):
        set_cell(tbl.cell(0, j), h, bold=True)
    for i, (det, name) in enumerate(sets, 1):
        set_cell(tbl.cell(i, 0), name, size=Pt(10))
        for j, sp in enumerate(("All", "D", "T"), 1):
            r = r1(sp, det)
            if r is None:
                set_cell(tbl.cell(i, j), "—", size=Pt(10))
                continue
            if det == "baseline_majority":
                set_cell(tbl.cell(i, j), f"{r.acc_mean:.3f}", size=Pt(10))
                continue
            sig = (pd.notna(r.perm_p) and float(r.perm_p) < 0.05
                   and pd.notna(r.p_ttest_maj) and float(r.p_ttest_maj) < 0.05)
            set_cell(tbl.cell(i, j), f"{r.acc_mean:.3f}", size=Pt(10), bold=sig)


s18 = find_slide("Detector vs 3 Priors")
title_sh = next(sh for sh in s18.shapes if sh.has_text_frame)
keep = title_sh._element
for sh in list(s18.shapes):
    if sh._element is not keep:
        sh._element.getparent().remove(sh._element)
title_sh.text_frame.text = "Detector vs Priors — Singles & Fusions (group-level accuracy)"
for run in title_sh.text_frame.paragraphs[0].runs:
    run.font.size = Pt(24)
    run.font.bold = True

add_matrix(s18, 0.35, "Single feature sets", SINGLES)
add_matrix(s18, 6.85, "Fusions", FUSIONS)

nb = s18.shapes.add_textbox(Inches(0.35), Inches(4.6), Inches(12.6), Inches(1.6))
nb.text_frame.word_wrap = True
lines = [
    "Bold = significant under BOTH tests (permutation & one-sample t vs majority, p<.05). "
    "185 group-windows, clean 2-annotator labels, nested LOSO.",
    "Text carries the signal: linguistic best single (All/Triads); ling+GxS best fusion (All 0.786, "
    "Dyads 0.842 — the only Dyads detector significant under both tests); Triads best = audio+linguistic (gaze-free).",
    "AIxVR is dominated everywhere, incl. its new linguistic fusion (ling+AIxVR ≤ linguistic alone in every split).",
    "Full stats (±SD, Cohen's d, exact p) per set: c2_paper_table.csv; effect-size tables on slides 12–13 & 20.",
]
for i, t in enumerate(lines):
    p = nb.text_frame.paragraphs[0] if i == 0 else nb.text_frame.add_paragraph()
    p.text = t
    for run in p.runs:
        run.font.size = Pt(10)

prs.save(SRC)
print(f"saved {SRC}: slide 18 rebuilt with two side-by-side tables")
