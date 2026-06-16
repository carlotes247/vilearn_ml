"""Append an openSMILE slide: raw interpretable acoustics (88 eGeMAPS functionals)
vs v/a/d affect. Addresses the 'v/a/d is only an estimate, raw acoustics should
count' point. Result: raw acoustics do NOT beat the affect estimate. Idempotent."""
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt

SRC = "discover/2026_March_ViLearn_Planning.pptx"
TITLE = "Raw Acoustics (openSMILE) vs Affect (v/a/d)"
prs = Presentation(SRC)
if any(any(sh.has_text_frame and sh.text_frame.text.strip() == TITLE for sh in s.shapes) for s in prs.slides):
    print("slide already present"); raise SystemExit

r = pd.read_csv("scratch/ideas/idea2_opensmile_results.csv")
piv = r.pivot_table(index=["target", "split"], columns="features", values="acc", aggfunc="first")

s = prs.slides.add_slide(prs.slide_layouts[7])
s.shapes.title.text = TITLE
header = ["Target", "Split", "v/a/d (3)", "openSMILE (88)", "v/a/d + openSMILE"]
rows = []
for (tgt, sp) in [("Group TE", "all"), ("Group TE", "D"), ("Group TE", "T"),
                  ("Individual", "all"), ("Individual", "D"), ("Individual", "T")]:
    row = piv.loc[(tgt, sp)]
    rows.append([tgt, sp, f"{row['vad']:.2f}", f"{row['opensmile']:.2f}", f"{row['vad+os']:.2f}"])
widths = [2.4, 1.2, 2.2, 2.6, 3.0]
tbl = s.shapes.add_table(len(rows) + 1, len(header), Inches(0.5), Inches(1.4),
                         Inches(sum(widths)), Inches(0.4 * (len(rows) + 1))).table
for j, wd in enumerate(widths):
    tbl.columns[j].width = Inches(wd)
for j, h in enumerate(header):
    c = tbl.cell(0, j); c.text = h
    c.text_frame.paragraphs[0].runs[0].font.size = Pt(12); c.text_frame.paragraphs[0].runs[0].font.bold = True
for i, row in enumerate(rows, 1):
    for j, v in enumerate(row):
        c = tbl.cell(i, j); c.text = str(v)
        c.text_frame.paragraphs[0].runs[0].font.size = Pt(11)

tb = s.shapes.add_textbox(Inches(0.5), Inches(4.6), Inches(12.0), Inches(2.3))
tb.text_frame.word_wrap = True
bullets = [
    "Addresses the objection that v/a/d is affect ESTIMATED from audio, not a raw acoustic feature.",
    "openSMILE = 88 interpretable eGeMAPS functionals (F0, loudness, jitter, shimmer, MFCC, spectral flux ...) — raw acoustics, not embeddings.",
    "Result: raw acoustics do NOT beat the affect estimate for group TE (v/a/d 0.69–0.72 > openSMILE 0.57–0.62 across all splits); combining them does not help.",
    "Only individual-All sees a small gain from adding openSMILE (0.64 vs 0.61). 88 features on ~20 sessions also risks overfitting.",
    "Takeaway: the compact 3-dim affect (v/a/d) is the stronger, more parsimonious audio representation — even raw acoustics don't outperform it.",
]
for i, t in enumerate(bullets):
    p = tb.text_frame.paragraphs[0] if i == 0 else tb.text_frame.add_paragraph()
    p.text = t
    for run in p.runs:
        run.font.size = Pt(12)

prs.save(SRC)
print(f"openSMILE slide added; {len(prs.slides)} slides")
