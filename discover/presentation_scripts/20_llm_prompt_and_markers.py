"""Add two backup slides after the LLM-rubric results slide:
(A) the exact prompt used to operationalize the TE schema (rubric + 4 axes + JSON spec);
(B) worked marker examples — Low/Med/High windows with their per-axis scores + label.
Idempotent. Slides appended then reordered (orphan-free) to sit after the rubric slide."""
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt

SRC = "discover/2026_March_ViLearn_Planning.pptx"
T_PROMPT = "LLM Rubric Prompt (operationalizing the TE schema)"
T_EX = "LLM Markers — Worked Examples (Low / Medium / High)"
prs = Presentation(SRC)
titles = {sh.text_frame.text.strip() for s in prs.slides for sh in s.shapes if sh.has_text_frame}
if T_PROMPT in titles and T_EX in titles:
    print("slides already present"); raise SystemExit

# locate the rubric results slide to anchor placement after it
anchor = next((i for i, s in enumerate(prs.slides)
               if any(sh.has_text_frame and sh.text_frame.text.strip().startswith("Operationalizing the TE Rubric")
                      for sh in s.shapes)), len(prs.slides) - 1)

# ---------- (A) prompt slide ----------
sp = prs.slides.add_slide(prs.slide_layouts[7])
sp.shapes.title.text = T_PROMPT
tb = sp.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(12.3), Inches(5.6))
tb.text_frame.word_wrap = True
blocks = [
    ("SYSTEM (rubric, verbatim from the study's coding schema):", True),
    ("Rate the TASK ENGAGEMENT of a group's discussion in a 60 s window of a collaborative VR "
     "learning task (German transcript, speakers tagged [blue]/[green]/[red]).", False),
    ("LOW — silence / no cognitive interaction; OR talk WITHOUT content engagement (off-topic).", False),
    ("MEDIUM — task discussed SUPERFICIALLY: repeating the task, opinions/experiences without "
     "reasoning, depth, or connections.", False),
    ("HIGH — ELABORATE: prior knowledge, making connections, hypotheticals, out-of-the-box.", False),
    ("Rate 0.0–1.0 on 4 axes:  task_relevance · content_depth · reasoning_present · connecting_ideas. "
     "Then te_continuous (0–1) and predicted_level (low/medium/high).", True),
    ('Output: JSON only — {"task_relevance":_, "content_depth":_, "reasoning_present":_, '
     '"connecting_ideas":_, "te_continuous":_, "predicted_level":"low|medium|high"}', False),
    ("USER:  Transcript:\\n[blue] … [green] … [red] …   (the window's diarized turns, time-ordered)", True),
    ("Model: local Qwen3.6-27B, temperature per HF (think: 1.0 / non-think: 0.7), seed 42, JSON-forced. "
     "One call per window (~185). Binarize te_continuous at 0.5 for the high/low prediction.", False),
]
for i, (txt, bold) in enumerate(blocks):
    p = tb.text_frame.paragraphs[0] if i == 0 else tb.text_frame.add_paragraph()
    p.text = txt
    for run in p.runs:
        run.font.size = Pt(12 if bold else 11); run.font.bold = bold

# ---------- (B) worked-examples slide ----------
se = prs.slides.add_slide(prs.slide_layouts[7])
se.shapes.title.text = T_EX
txt = pd.read_csv("scratch/ideas/idea3_window_text.csv")
sc = pd.read_csv("scratch/ideas/idea3_window_scores_think.csv")
m = txt.merge(sc, on=["group_name", "sec"], suffixes=("", "_s"))


def pick(level, key):
    sub = m[m.predicted_level == level].sort_values("te_continuous")
    return sub.iloc[len(sub) // 2] if key == "mid" else (sub.iloc[1] if key == "lo" else sub.iloc[-2])


rows_src = [("Low", pick("low", "lo")), ("Medium", pick("medium", "mid")), ("High", pick("high", "hi"))]
header = ["Level", "Transcript (excerpt)", "task_rel", "depth", "reason", "connect", "te", "y"]
widths = [0.9, 5.0, 1.0, 0.9, 0.9, 1.0, 0.7, 0.5]
tbl = se.shapes.add_table(4, len(header), Inches(0.3), Inches(1.3),
                          Inches(sum(widths)), Inches(1.4)).table
for j, wd in enumerate(widths):
    tbl.columns[j].width = Inches(wd)
for j, h in enumerate(header):
    c = tbl.cell(0, j); c.text = h
    c.text_frame.paragraphs[0].runs[0].font.size = Pt(10); c.text_frame.paragraphs[0].runs[0].font.bold = True
for i, (lvl, r) in enumerate(rows_src, 1):
    ex = " ".join(r.text.replace("\n", " ").split())[:150] + "…"
    cells = [lvl, ex, f"{r.task_relevance:.2f}", f"{r.content_depth:.2f}",
             f"{r.reasoning_present:.2f}", f"{r.connecting_ideas:.2f}",
             f"{r.te_continuous:.2f}", str(int(r.y))]
    for j, v in enumerate(cells):
        c = tbl.cell(i, j); c.text = v
        c.text_frame.paragraphs[0].runs[0].font.size = Pt(9)

nb = se.shapes.add_textbox(Inches(0.3), Inches(3.2), Inches(12.5), Inches(2.0))
nb.text_frame.word_wrap = True
for i, t in enumerate([
    "The 4 axes ARE the interpretable output: e.g. Low has task_relevance≈0.1 (off-task hand-gesture chatter); High scores high on all four (reasoning + connecting prior knowledge).",
    "These per-axis scores are the features fed to the panel (idea (b)); te_continuous>0.5 is the zero-shot prediction (idea (a)).",
    "Medium is the ambiguous band where annotators disagree (weak κ) — kept here, binarized at 0.5.",
]):
    p = nb.text_frame.paragraphs[0] if i == 0 else nb.text_frame.add_paragraph()
    p.text = "• " + t
    for run in p.runs:
        run.font.size = Pt(11)

# reorder: move the two new slides (last two) to just after the anchor
xs = prs.slides._sldIdLst
els = list(xs)
new_prompt, new_ex = els[-2], els[-1]
xs.remove(new_prompt); xs.remove(new_ex)
xs.insert(anchor + 1, new_prompt)
xs.insert(anchor + 2, new_ex)
prs.save(SRC)
print(f"prompt + markers slides added after slide index {anchor}; {len(prs.slides)} slides")
