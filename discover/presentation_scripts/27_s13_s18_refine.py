"""Slide 13 + 18 refinements (saveli deck review round 2, 2026-07-07).

Slide 13: 'best single' -> 'best single set' (GxS is cross-modal gaze×speaking, so
'single modality' would be wrong); Dyads duplicate rows merged into
'prior gaze (GxS) = best single set'; clarifying bullet appended.

Slide 18 (two-table layout kept, rebuilt): +'#f' column (per-split counts like 2/2/1
for AIxVR sets), no 'gaze' suffix in labels, cells = 'acc (d)' with significance
markers: * = permutation p<.05, † = t-test-vs-majority p<.05.

Run from repo root:  python3 discover/presentation_scripts/27_s13_s18_refine.py
"""
import pandas as pd
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

SRC = "scratch/2026_March_ViLearn_Planning.pptx"
prs = Presentation(SRC)
c2 = pd.read_csv("scratch/te_adoption/a1_out/c2_paper_table.csv")


def r1(sp, det):
    m = c2[(c2.split == sp) & (c2.detector == det) & (c2["rank"] == 1)]
    return m.iloc[0] if len(m) else None


def find_slide(prefix):
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip().startswith(prefix):
                return s
    return None


def set_cell(cell, text, size=Pt(10), bold=False):
    cell.text = str(text)
    r = cell.text_frame.paragraphs[0].runs[0]
    r.font.size = size
    r.font.bold = bold


# ---------------- slide 13 ----------------
s13 = find_slide("Main Results — Group-Level")
tbl = next(sh.table for sh in s13.shapes if sh.has_table)
dup_tr = None
for row in tbl.rows:
    lab = row.cells[1].text.strip()
    if lab.startswith("best single:"):
        sp_lab = row.cells[0].text.strip()
        if sp_lab == "Dyads" and "GazexSpeaking" in lab:
            dup_tr = row._tr  # duplicate of the prior-gaze row -> drop
        else:
            set_cell(row.cells[1], lab.replace("best single:", "best single set:"), size=Pt(9))
    elif lab == "prior gaze (GxS)" and row.cells[0].text.strip() == "Dyads":
        set_cell(row.cells[1], "prior gaze (GxS) = best single set", size=Pt(9))
    elif lab.startswith("best fusion:"):
        set_cell(row.cells[1], lab.replace("best fusion:", "best fused set:"), size=Pt(9))
if dup_tr is not None:
    dup_tr.getparent().remove(dup_tr)
    print("slide 13: Dyads duplicate row removed")
# clarifying bullet
for sh in s13.shapes:
    if sh.has_text_frame and "Both significance tests" in sh.text_frame.text:
        if "single set ≠ single modality" not in sh.text_frame.text:
            p = sh.text_frame.add_paragraph()
            p.text = ("Terminology: 'set' not 'modality' — GazexSpeaking is itself cross-modal "
                      "(gaze × speaking status), so rows compare feature SETS; single set ≠ single modality.")
            for run in p.runs:
                run.font.size = Pt(10)
        break
print("slide 13: relabeled to single/fused SET")

# ---------------- slide 18 ----------------
SINGLES = [("baseline_majority", "majority baseline"), ("AIxVR", "AIxVR"),
           ("GazexSpeaking", "GazexSpeaking"), ("audio", "audio (v/a/d)"),
           ("linguistic", "linguistic")]
FUSIONS = [("audio+AIxVR", "audio+AIxVR"), ("audio+GazexSpeaking", "audio+GxS"),
           ("audio+linguistic", "audio+linguistic"), ("linguistic+AIxVR", "ling+AIxVR"),
           ("linguistic+GazexSpeaking", "ling+GxS"),
           ("audio+linguistic+GazexSpeaking", "audio+ling+GxS")]


def nfeat_label(det):
    ns = [int(r1(sp, det).n_feat) for sp in ("All", "D", "T") if r1(sp, det) is not None]
    if not ns:
        return "—"
    return str(ns[0]) if len(set(ns)) == 1 else "/".join(map(str, ns))


def cell_text(r, baseline=False):
    if baseline:
        return f"{r.acc_mean:.3f}"
    marks = ""
    if pd.notna(r.perm_p) and str(r.perm_p) != "" and float(r.perm_p) < 0.05:
        marks += "*"
    if pd.notna(r.p_ttest_maj) and float(r.p_ttest_maj) < 0.05:
        marks += "†"
    return f"{r.acc_mean:.3f} ({r.cohen_d_maj:.2f}){marks}"


def add_matrix(slide, x, title, sets):
    tb = slide.shapes.add_textbox(Inches(x), Inches(1.05), Inches(6.0), Inches(0.4))
    tb.text_frame.paragraphs[0].text = title
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    tb.text_frame.paragraphs[0].runs[0].font.bold = True
    widths = [1.75, 0.5, 1.35, 1.35, 1.35]
    tbl = slide.shapes.add_table(len(sets) + 1, 5, Inches(x), Inches(1.5),
                                 Inches(sum(widths)), Inches(0.34 * (len(sets) + 1))).table
    for j, wd in enumerate(widths):
        tbl.columns[j].width = Inches(wd)
    for j, h in enumerate(["Feature set", "#f", "All", "Dyads", "Triads"]):
        set_cell(tbl.cell(0, j), h, bold=True)
    for i, (det, name) in enumerate(sets, 1):
        set_cell(tbl.cell(i, 0), name, size=Pt(9))
        set_cell(tbl.cell(i, 1), "—" if det == "baseline_majority" else nfeat_label(det), size=Pt(9))
        for j, sp in enumerate(("All", "D", "T"), 2):
            r = r1(sp, det)
            set_cell(tbl.cell(i, j), "—" if r is None else cell_text(r, det == "baseline_majority"),
                     size=Pt(9))


s18 = find_slide("Detector vs Priors")
title_sh = next(sh for sh in s18.shapes if sh.has_text_frame)
keep = title_sh._element
for sh in list(s18.shapes):
    if sh._element is not keep:
        sh._element.getparent().remove(sh._element)
title_sh.text_frame.text = "Detector vs Priors — Singles & Fusions (acc, effect size, both tests)"
for run in title_sh.text_frame.paragraphs[0].runs:
    run.font.size = Pt(22)
    run.font.bold = True

add_matrix(s18, 0.25, "Single feature sets", SINGLES)
add_matrix(s18, 6.85, "Fusions", FUSIONS)

nb = s18.shapes.add_textbox(Inches(0.25), Inches(4.7), Inches(12.7), Inches(1.7))
nb.text_frame.word_wrap = True
lines = [
    "Cell = accuracy (Cohen's d vs majority floor); * = permutation p<.05, † = one-sample t-test vs "
    "majority p<.05. #f = features (All/Dyads/Triads where split-specific). 185 group-windows, clean "
    "2-annotator labels, nested LOSO.",
    "Text carries the signal: linguistic best single set (All/Triads); ling+GxS best fusion "
    "(Dyads 0.842 — only Dyads detector with both markers); Triads best = audio+linguistic (gaze-free).",
    "AIxVR dominated everywhere incl. ling+AIxVR ≤ linguistic alone. Full stats: c2_paper_table.csv.",
]
for i, t in enumerate(lines):
    p = nb.text_frame.paragraphs[0] if i == 0 else nb.text_frame.add_paragraph()
    p.text = t
    for run in p.runs:
        run.font.size = Pt(10)

prs.save(SRC)
print(f"saved {SRC}")
