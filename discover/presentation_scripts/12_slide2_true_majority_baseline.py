"""Slide 2 (prior GazexSpeaking comparison): replace the stochastic uniform/seed
dummy baseline (0.53/0.43/0.45) with the TRUE deterministic majority baseline
(always-predict-majority, per split), computed on the prior's own labels:
Dyads 0.62 / Triads 0.55 / All 0.52. Deterministic -> no SD. Add a footnote.
Idempotent. Run from repo root."""
from pptx import Presentation
from pptx.util import Inches, Pt

SRC = "discover/2026_March_ViLearn_Planning.pptx"
prs = Presentation(SRC)
s = prs.slides[1]  # slide #2

# row 2 = baseline; cols: 0 label,1 D-M,2 D-SD,3 T-M,4 T-SD,5 All-M,6 All-SD
NEW = {0: "Majority baseline", 1: "0.62", 2: "—", 3: "0.55", 4: "—", 5: "0.52", 6: "—"}


def set_cell(cell, text, sz=None, bold=None):
    old = cell.text_frame.paragraphs[0].runs[0] if cell.text_frame.paragraphs[0].runs else None
    sz = sz or (old.font.size if old and old.font.size else Pt(12))
    bold = bold if bold is not None else (old.font.bold if old else False)
    cell.text = str(text)
    r = cell.text_frame.paragraphs[0].runs[0]
    r.font.size = sz
    r.font.bold = bool(bold)


for sh in s.shapes:
    if sh.has_table and any("GazexSpeaking" in c.text for c in sh.table.rows[0].cells):
        t = sh.table
        # find the baseline row by its first cell
        for ri in range(len(t.rows)):
            if t.cell(ri, 0).text.strip().lower().startswith("baseline") or \
               t.cell(ri, 0).text.strip() == "Majority baseline":
                for cj, val in NEW.items():
                    set_cell(t.cell(ri, cj), val)
                break

# footnote (idempotent)
if not any(sh.has_text_frame and "Majority baseline =" in sh.text_frame.text for sh in s.shapes):
    tb = s.shapes.add_textbox(Inches(0.4), Inches(6.6), Inches(12.4), Inches(0.9))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    p.text = ("Majority baseline = always-predict-majority (deterministic, per split, on the prior's "
              "labels). Replaces the original stochastic uniform/stratified dummy (0.53/0.43/0.45), which "
              "was a weaker floor — the models' margin over the honest baseline is much smaller "
              "(e.g. SVM Dyads 0.66 vs 0.62; SVM Triads 0.63 vs 0.55).")
    p.runs[0].font.size = Pt(10)
    p.runs[0].font.italic = True

prs.save(SRC)
print("slide 2 baseline -> true majority (0.62/0.55/0.52) + footnote")
