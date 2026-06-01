from pptx import Presentation
from pptx.util import Inches, Pt

SRC = "discover/2026_March_ViLearn_Planning.pptx"
prs = Presentation(SRC)


def find_idx(substr):
    for i, s in enumerate(prs.slides):
        for sh in s.shapes:
            if sh.has_text_frame and substr.lower() in sh.text_frame.text.lower():
                return i
    return None


# --- 1) Rewrite the RQ slide body with concrete answers ---
rq_idx = find_idx("Answers to Research Questions")
s = prs.slides[rq_idx]
body = None
for ph in s.placeholders:
    if ph.placeholder_format.type != 1:  # not TITLE
        body = ph
tf = body.text_frame
tf.clear()
bullets = [
    "Double dissociation: group TE <- audio affect (valence/dominance); individual engagement <- (para)linguistic (speaking + text)",
    "Linguistic-only detects TE? Modest for group TE (acc ~0.72, below audio); but the BEST channel for individual engagement (acc ~0.71)",
    "Strongest linguistic features? speaking_role, word_count, words_per_second + sentiment (individual). Group TE driven by audio valence (+) / dominance (-), not text",
    "Multimodal (both)? group TE 0.82-0.88, individual 0.69-0.73 = matches the dominant single modality; no synergy; embeddings on top degrade slightly",
    "Multimodal significantly beats prior gaze x speaking QDA on All + Triads (d = 0.9-1.35); dyads trend same but n=8 underpowered",
    "Open: individual engagement ~0.70; class imbalance (74% high); 60 s windows best granularity",
    "1-year: real-time multimodal per-window TE detector feeding adaptive VR feedback",
]
for i, t in enumerate(bullets):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.text = t
    for r in p.runs:
        r.font.size = Pt(14)

# --- 2) Add an ablation results slide right after the Feature Modalities slide ---
slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
slide.shapes.title.text = "Modality Ablation — what each modality detects (60 s, LOSO)"
header = ["Target (split)", "Multimodal", "Audio-only", "(Para)ling-only"]
rows = [
    ["Group TE (All)", "0.82", "0.82", "0.72"],
    ["Group TE (Triads)", "0.88", "0.86", "0.72"],
    ["Individual (All)", "0.69", "0.66", "0.71"],
    ["Individual (Triads)", "0.73", "0.64", "0.72"],
]
left, top, width = Inches(0.5), Inches(1.6), Inches(10.0)
tbl = slide.shapes.add_table(len(rows) + 1, 4, left, top, width, Inches(0.4 * (len(rows) + 1))).table
for j, w in enumerate([3.4, 2.2, 2.2, 2.2]):
    tbl.columns[j].width = Inches(w)
for j, h in enumerate(header):
    c = tbl.cell(0, j); c.text = h
    r = c.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(13); r.font.bold = True
for i, row in enumerate(rows, 1):
    for j, val in enumerate(row):
        c = tbl.cell(i, j); c.text = str(val)
        c.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
note = ("Best model per cell, accuracy. DOUBLE DISSOCIATION: group task engagement is carried by "
        "audio affect (valence/dominance) - text adds nothing; individual engagement is carried by "
        "(para)linguistic (speaking + word rate) - group-shared audio gives no per-role signal. "
        "Combining modalities = the dominant one, no synergy.")
tb = slide.shapes.add_textbox(left, top + Inches(0.4 * (len(rows) + 1)) + Inches(0.2), Inches(11.5), Inches(1.4))
tb.text_frame.word_wrap = True
p = tb.text_frame.paragraphs[0]; p.text = note
p.runs[0].font.size = Pt(12); p.runs[0].font.italic = True

# move new slide (last) to just after the Feature Modalities slide
fm_idx = find_idx("Feature Modalities")
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
sldIdLst.remove(ids[-1])
sldIdLst.insert(fm_idx + 1, ids[-1])

prs.save(SRC)
print(f"RQ slide idx {rq_idx} rewritten; ablation slide inserted after Feature Modalities (idx {fm_idx}); {len(prs.slides)} slides")
