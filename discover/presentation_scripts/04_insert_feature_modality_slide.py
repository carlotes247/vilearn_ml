from pptx import Presentation
from pptx.util import Inches, Pt

SRC = "discover/2026_March_ViLearn_Planning.pptx"
prs = Presentation(SRC)
slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
slide.shapes.title.text = "Feature Modalities (ablation buckets)"

header = ["Modality", "Features"]
rows = [
    ["(Para)linguistic", "text: word_count, words/sec, avg_word_length, question/statement_rate, "
                         "sentiment;  speaking: speaking_role, speaking_seconds, segment_count"],
    ["Audio (affect)", "arousal, dominance, valence  (group-level audio emotion)"],
    ["Multimodal", "(para)linguistic + audio  (= our main detector)"],
    ["(+ embeddings)", "optional: openSMILE + emoW2V + sentiment embeddings, PCA (degrade slightly)"],
]
left, top, width = Inches(0.5), Inches(1.6), Inches(12.2)
tbl = slide.shapes.add_table(len(rows) + 1, 2, left, top, width, Inches(0.4 * (len(rows) + 1))).table
tbl.columns[0].width = Inches(2.6)
tbl.columns[1].width = Inches(9.6)
for j, h in enumerate(header):
    c = tbl.cell(0, j); c.text = h
    r = c.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(13); r.font.bold = True
for i, row in enumerate(rows, 1):
    for j, val in enumerate(row):
        c = tbl.cell(i, j); c.text = str(val)
        c.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
        c.text_frame.word_wrap = True

note = ("Audio is group-shared (one group.audio.wav); the transcript is diarized from it. "
        "Speaking attributes the shared signal to each role -> folded into (para)linguistic. "
        "v/a/d derived from audio -> audio. Ablation: (para)linguistic-only vs audio-only vs both.")
tb = slide.shapes.add_textbox(left, top + Inches(0.4 * (len(rows) + 1)) + Inches(0.2), width, Inches(1.0))
tb.text_frame.word_wrap = True
p = tb.text_frame.paragraphs[0]; p.text = note
p.runs[0].font.size = Pt(11); p.runs[0].font.italic = True

# Move the new slide (currently last) to position 10 (after the Setup slide, idx 9)
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
sldIdLst.remove(ids[-1])
sldIdLst.insert(10, ids[-1])

prs.save(SRC)
print("inserted feature-modality slide at position 11;", len(prs.slides), "slides")
