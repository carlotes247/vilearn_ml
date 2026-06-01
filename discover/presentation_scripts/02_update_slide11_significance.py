from pptx import Presentation
from pptx.util import Inches, Pt

SRC = "discover/2026_March_ViLearn_Planning.pptx"
prs = Presentation(SRC)
slide = prs.slides[10]  # slide 11 (0-indexed)

# Remove every shape except the title placeholder (old table + note textbox).
title = slide.shapes.title
for sh in list(slide.shapes):
    if sh is title:
        continue
    sh._element.getparent().remove(sh._element)

title.text = "Group TE Detection: Results + Significance (vs prior detector)"

header = ["Split (folds)", "Best model", "Acc (mean±SD)", "F1", "vs baseline", "vs prior GazexSpeak QDA"]
rows = [
    ["Dyads (8)", "QDA", "0.74 ± 0.21", "0.57", "ns (d=+0.89)", "ns (d=+0.52)"],
    ["Triads (12)", "Random Forest", "0.88 ± 0.09", "0.80", "p<0.001  d=2.4 ✓", "p=0.005  d=1.35 ✓"],
    ["All (20)", "Random Forest", "0.82 ± 0.16", "0.71", "p<0.001  d=1.48 ✓", "p=0.004  d=0.92 ✓"],
]
left, top, width = Inches(0.4), Inches(1.5), Inches(11.4)
height = Inches(0.35 * (len(rows) + 1))
tbl = slide.shapes.add_table(len(rows) + 1, len(header), left, top, width, height).table
widths = [1.5, 2.0, 1.9, 0.8, 2.4, 2.8]
for j, w in enumerate(widths):
    tbl.columns[j].width = Inches(w)
for j, h in enumerate(header):
    c = tbl.cell(0, j)
    c.text = h
    r = c.text_frame.paragraphs[0].runs[0]
    r.font.size = Pt(12); r.font.bold = True
for i, row in enumerate(rows, 1):
    for j, val in enumerate(row):
        c = tbl.cell(i, j)
        c.text = str(val)
        c.text_frame.paragraphs[0].runs[0].font.size = Pt(12)

note = ("Per-fold leave-one-group-out accuracy; paired t-test (Bonferroni) + Cohen's d. "
        "Linguistic+affective beats the prior gaze×speaking QDA on Triads & All "
        "(medium–large effect, d 0.9–1.35); dyads trend the same way but n=8 is underpowered. "
        "No gaze features used.")
tb = slide.shapes.add_textbox(left, top + height + Inches(0.2), width, Inches(1.2))
tf = tb.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = note
p.runs[0].font.size = Pt(12)
p.runs[0].font.italic = True

prs.save(SRC)
print("updated slide 11;", len(prs.slides), "slides total")
