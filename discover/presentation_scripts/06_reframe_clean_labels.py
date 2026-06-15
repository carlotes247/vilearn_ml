"""Reframe deck slides 12-17 for HARMONIZED 2-annotator labels + A1 fair comparison.

Idempotent in-place patcher (matches tables/notes by content, overwrites values).
Backs up to *_preclean_backup.pptx. Run from repo root. CURRENT deck indices
(post-04/05 insertion): 11=Modality Ablation, 12=Results vs prior, 13=Granularity,
14=Feature contribution, 15=Embeddings, 16=RQ.
"""
import os
import shutil
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import PP_PLACEHOLDER

SRC = "discover/2026_March_ViLearn_Planning.pptx"
BACKUP = SRC.replace(".pptx", "_preclean_backup.pptx")
if not os.path.exists(BACKUP):  # preserve the pristine pre-reframe deck across re-runs
    shutil.copyfile(SRC, BACKUP)
prs = Presentation(SRC)

# ---- embeddings numbers (clean labels, from compute_extra_panel_numbers.py) ----
# rows: model -> (base, +embeddings, delta)  -- FILLED from compute output
EMBED = {
    "Random Forest": ("0.67", "0.63", "-0.04"),
    "Logistic Regression": ("0.68", "0.55", "-0.13"),
    "Naive Bayes": ("0.62", "0.58", "-0.04"),
}


def set_cell(cell, text, sz=12, bold=False):
    cell.text = str(text)
    r = cell.text_frame.paragraphs[0].runs[0]
    r.font.size = Pt(sz)
    r.font.bold = bold


def find_table(slide, header_sig):
    for sh in slide.shapes:
        if sh.has_table:
            hdr = " ".join(c.text for c in sh.table.rows[0].cells)
            if header_sig in hdr:
                return sh.table
    return None


def replace_note(slide, contains, new_text):
    for sh in slide.shapes:
        if sh.has_text_frame and contains in sh.text_frame.text:
            tf = sh.text_frame
            # keep first paragraph, set its text, drop the rest
            for p in tf.paragraphs[1:]:
                p._p.getparent().remove(p._p)
            p0 = tf.paragraphs[0]
            for r in list(p0.runs)[1:]:
                r._r.getparent().remove(r._r)
            if p0.runs:
                p0.runs[0].text = new_text
            else:
                p0.text = new_text
            return True
    return False


# ===== Slide 11 (#12): Modality Ablation — clean best-per-cell =====
t = find_table(prs.slides[11], "Multimodal")
abl = {
    "Group TE (All)": ("0.68", "0.68", "0.64"),
    "Group TE (Triads)": ("0.70", "0.70", "0.62"),
    "Group TE (Dyads)": ("0.65", "0.71", "0.57"),
    "Individual (All)": ("0.70", "0.64", "0.70"),
    "Individual (Triads)": ("0.73", "0.65", "0.72"),
    "Individual (Dyads)": ("0.68", "0.60", "0.66"),
}
for ri in range(1, len(t.rows)):
    label = t.cell(ri, 0).text.strip()
    if label in abl:
        for cj, v in zip((1, 2, 3), abl[label]):
            set_cell(t.cell(ri, cj), v)

# ===== Slide 12 (#13): Group TE Results + Significance (vs prior) =====
t = find_table(prs.slides[12], "Split (folds)")
res = [
    ["Dyads (8)", "Decision Tree", "0.65 ± 0.11", "0.54", "n.s. (d=0.63)", "n.s. (his 0.77; d=−0.58)"],
    ["Triads (12)", "QDA", "0.70 ± 0.20", "0.63", "n.s. (d=0.54)", "n.s. (his 0.65; d=+0.36)"],
    ["All (20)", "Logistic Regression", "0.68 ± 0.18", "0.57", "n.s. (d=0.58)", "n.s. (his 0.73; d=−0.36)"],
]
for ri, row in enumerate(res, 1):
    for cj, v in enumerate(row):
        set_cell(t.cell(ri, cj), v)
replace_note(prs.slides[12], "Per-fold",
             "Per-fold LOSO accuracy; paired t-test (Bonferroni) + Cohen's d. With harmonized "
             "2-annotator labels + identical nested CV, our multimodal detector and the prior gaze×speaking "
             "QDA are statistically indistinguishable on every split (all n.s.) — our detector edges triads, "
             "gaze edges dyads. Nothing survives Bonferroni vs chance (n=8–20 underpowered). The earlier "
             "'beats prior' result was a label-distribution artifact.")

# ===== Slide 13 (#14): Granularity — drop 1 Hz row, clean numbers =====
s = prs.slides[13]
t = find_table(s, "Granularity")
# rebuild as 2 data rows (header + 2)
old = t._tbl
# easiest: overwrite first two data rows, delete any extra
rows_data = [["Transcript segment", "0.55", "per-utterance, noisy"],
             ["60 s window", "0.68", "best — matches prior work"]]
# delete extra data rows beyond what we need (keep header + 2)
while len(t.rows) > 1 + len(rows_data):
    last = t.rows[len(t.rows) - 1]._tr
    last.getparent().remove(last)
for ri, row in enumerate(rows_data, 1):
    for cj, v in enumerate(row):
        set_cell(t.cell(ri, cj), v)
replace_note(s, "granularity",
             "Best model per granularity (group TE, LOSO). Coarser aggregation → less noise; 60 s wins. "
             "(1 Hz frame omitted — not recomputed on harmonized labels.)")

# ===== Slide 14 (#15): Feature contribution — clean coefs =====
t = find_table(prs.slides[14], "LogReg coef")
feats = [["dominance", "-1.43", "0.18"],
         ["valence", "+0.58", "0.16"],
         ["question_rate", "-0.35", "0.04"],
         ["arousal", "+0.19", "0.12"]]
while len(t.rows) > 1 + len(feats):
    last = t.rows[len(t.rows) - 1]._tr
    last.getparent().remove(last)
for ri, row in enumerate(feats, 1):
    for cj, v in enumerate(row):
        set_cell(t.cell(ri, cj), v)
replace_note(prs.slides[14], "driven by",
             "Group TE driven by audio affect (dominance −, valence +); higher question rate → lower TE. "
             "Individual engagement: speaking activity + word-rate + text.")

# ===== Slide 15 (#16): Embeddings — clean base-vs-embeddings =====
t = find_table(prs.slides[15], "Base acc")
for ri in range(1, len(t.rows)):
    label = t.cell(ri, 0).text.strip()
    if label in EMBED:
        for cj, v in zip((1, 2, 3), EMBED[label]):
            set_cell(t.cell(ri, cj), v)
replace_note(prs.slides[15], "added on top",
             "Embeddings (openSMILE + emoW2V + sentiment, PCA) added on top of base features. "
             "Net degradation — best base 0.68 > best +embeddings 0.63. Added dimensions = noise "
             "(harmonized 2-annotator labels).")

# ===== Slide 16 (#17): RQ answers — reframe to fair-tie =====
s = prs.slides[16]
body = None
for ph in s.placeholders:
    if ph.placeholder_format.type != PP_PLACEHOLDER.TITLE:
        body = ph
if body is not None:
    tf = body.text_frame
    tf.clear()
    bullets = [
        "Features = linguistic (text) + audio-derived affect (valence/arousal/dominance) + speaking → a MULTIMODAL detector, not linguistic-only",
        "With harmonized 2-annotator labels + nested CV: multimodal group TE ≈ 0.68 All / 0.70 Triads / 0.65 Dyads — statistically TIED with the prior gaze×speaking QDA (all n.s.); our detector edges triads, gaze edges dyads",
        "Double dissociation holds: group TE ← audio affect (valence/dominance); individual engagement ← (para)linguistic (speaking + text)",
        "Carlos's published baseline was itself depressed by a label time-stretch bug (clean QDA 0.67 vs published 0.60)",
        "Embeddings (openSMILE/emoW2V/sentiment, PCA) on top of base: no gain",
        "Open: nothing survives Bonferroni vs chance (n=8–20 underpowered); labels now balanced ~49% high (was 74% skew)",
        "1-year: real-time multimodal per-window TE detector feeding adaptive VR feedback",
    ]
    for i, text in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        for r in p.runs:
            r.font.size = Pt(15)

prs.save(SRC)
print("reframed slides 11,12,13,14,15,16;", len(prs.slides), "slides")
